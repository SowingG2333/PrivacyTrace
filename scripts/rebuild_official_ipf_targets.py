#!/usr/bin/env python3
"""Rebuild the non-official IPF targets from official statistical sources.

The script deliberately records transformations and proxy semantics.  In
particular, UN DESA's migrant-stock matrix is based on country of birth where
available and country of citizenship otherwise.  It is therefore used as a
direct observation for one of the two fields and as an official-data proxy for
the other, with the destination-specific basis retained in the audit output.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import json
import math
import re
import sys
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TARGET = PROJECT_ROOT / "data/targets/demographic_targets_2025.json"
DEFAULT_AUDIT = PROJECT_ROOT / "data/sources/official_ipf/derived_inputs_2025.json"
DEFAULT_VALUES_DOC = PROJECT_ROOT / "docs/current_ipf_constraint_values.md"
DEFAULT_CACHE = PROJECT_ROOT / ".cache/official_ipf"
DEFAULT_CITY_CATALOG = PROJECT_ROOT / "data/catalogs/wup2025_cities_2025.json"
DEFAULT_IHME_GBD_2023 = DEFAULT_CACHE / "ihme_gbd_2023.csv"

WPP_URL = (
    "https://population.un.org/wpp/assets/Excel%20Files/"
    "1_Indicator%20(Standard)/CSV_FILES/"
    "WPP2024_PopulationByAge5GroupSex_Medium.csv.gz"
)
IMS_URL = (
    "https://www.un.org/development/desa/pd/sites/"
    "www.un.org.development.desa.pd/files/"
    "undesa_pd_2020_ims_stock_by_sex_destination_and_origin.xlsx"
)
UNDATA_URL = (
    "https://data.un.org/Handlers/DownloadHandler.ashx?"
    "DataFilter=tableCode:{table}%3BcountryCode:{country_code}"
    "&DataMartId=POP&Format=csv"
)
PIP_URL = (
    "https://api.worldbank.org/pip/v1/pip?"
    "country={country}&year={year}&povline={line}"
    "&ppp_version=2021&format=json"
)
OECD_MIDDLE_INCOME_URL = (
    "https://www.oecd.org/en/publications/under-pressure-the-squeezed-middle-class_"
    "689afed1-en/full-report/component-6.html"
)
ILO_OCCUPATION_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_2EMP_SEX_OCU_NB_A.rds"
)
ILO_EMPLOYMENT_RATE_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_2WAP_SEX_AGE_RT_A.rds"
)
ILO_EMPLOYMENT_RATE_BY_AGE_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_DWAP_SEX_AGE_RT_A.rds"
)
ILO_WORKING_AGE_AGE_EDUCATION_URL = (
    "https://rplumber.ilo.org/files/indicator/POP_XWAP_SEX_AGE_EDU_NB_A.rds"
)
ILO_EMPLOYMENT_AGE_EDUCATION_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_TEMP_SEX_AGE_EDU_NB_A.rds"
)
ILO_EMPLOYMENT_RATE_BY_EDUCATION_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_DWAP_SEX_EDU_RT_A.rds"
)
ILO_OCCUPATION_EDUCATION_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_TEMP_SEX_OCU_EDU_NB_A.rds"
)
ILO_AGE_OCCUPATION_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_TEMP_SEX_AGE_OCU_NB_A.rds"
)
ILO_AGE_MARITAL_EMPLOYMENT_RATE_URL = (
    "https://rplumber.ilo.org/files/indicator/EMP_DWAP_SEX_AGE_MTS_RT_A.rds"
)
WORLD_BANK_WELFARE_PROFILE_URL = (
    "https://reproducibility.worldbank.org/catalog/250/download/732"
)
WORLD_BANK_WELFARE_PROFILE_MEMBER = (
    "FR_WLD_2024_198/Reproducibility package/Chapter 1/1-data/raw/other/"
    "global_profile_lineup_2022_country_level.dta"
)
WUP_CITY_URL = (
    "https://population.un.org/wup/assets/Download/Cities/"
    "WUP2025-F21-DEGURBA-Cities_Pop.xlsx"
)
PEW_RELIGION_URL = (
    "https://www.pewresearch.org/wp-content/uploads/sites/20/2025/06/"
    "Religious-Composition-2010-2020-dataset.zip"
)
WORLD_BANK_API = (
    "https://api.worldbank.org/v2/country/{countries}/indicator/{indicator}"
    "?format=json&per_page=20000"
)
COUNTRY_CODES = {
    "United States": {"m49": 840, "iso3": "USA", "un_name": "United States of America"},
    "China": {"m49": 156, "iso3": "CHN", "un_name": "China"},
    "India": {"m49": 356, "iso3": "IND", "un_name": "India"},
    "Brazil": {"m49": 76, "iso3": "BRA", "un_name": "Brazil"},
    "Mexico": {"m49": 484, "iso3": "MEX", "un_name": "Mexico"},
    "Nigeria": {"m49": 566, "iso3": "NGA", "un_name": "Nigeria"},
    "Germany": {"m49": 276, "iso3": "DEU", "un_name": "Germany"},
    "Japan": {"m49": 392, "iso3": "JPN", "un_name": "Japan"},
    "United Kingdom": {"m49": 826, "iso3": "GBR", "un_name": "United Kingdom"},
    "Canada": {"m49": 124, "iso3": "CAN", "un_name": "Canada"},
    "Colombia": {"m49": 170, "iso3": "COL", "un_name": "Colombia"},
    "France": {"m49": 250, "iso3": "FRA", "un_name": "France"},
    "Poland": {"m49": 616, "iso3": "POL", "un_name": "Poland"},
    "South Korea": {"m49": 410, "iso3": "KOR", "un_name": "Republic of Korea"},
    "Indonesia": {"m49": 360, "iso3": "IDN", "un_name": "Indonesia"},
    "Philippines": {"m49": 608, "iso3": "PHL", "un_name": "Philippines"},
    "Vietnam": {"m49": 704, "iso3": "VNM", "un_name": "Viet Nam"},
    "Pakistan": {"m49": 586, "iso3": "PAK", "un_name": "Pakistan"},
    "Bangladesh": {"m49": 50, "iso3": "BGD", "un_name": "Bangladesh"},
    "Egypt": {"m49": 818, "iso3": "EGY", "un_name": "Egypt"},
    "Turkey": {"m49": 792, "iso3": "TUR", "un_name": "Türkiye"},
    "Iran": {
        "m49": 364,
        "iso3": "IRN",
        "un_name": "Iran (Islamic Republic of)",
    },
    "Ethiopia": {"m49": 231, "iso3": "ETH", "un_name": "Ethiopia"},
    "South Africa": {"m49": 710, "iso3": "ZAF", "un_name": "South Africa"},
    "Kenya": {"m49": 404, "iso3": "KEN", "un_name": "Kenya"},
    "Australia": {"m49": 36, "iso3": "AUS", "un_name": "Australia"},
}

AGE_GROUPS = {
    "18-24": (18, 24),
    "25-34": (25, 34),
    "35-44": (35, 44),
    "45-54": (45, 54),
    "55-64": (55, 64),
    "65-80": (65, 80),
}
SEX_MAP = {"Male": "male", "Female": "female"}
RELATIONSHIP_MAP = {
    "Single (never married)": "single",
    "In consensual union": "in relationship",
    "Married": "married",
    "Divorced and not remarried": "divorced",
    "Married but separated": "divorced",
    "Widowed and not remarried": "widowed",
}
RELATIONSHIPS = [
    "single",
    "in relationship",
    "married",
    "divorced",
    "widowed",
]
ETHNICITIES = [
    "East Asian",
    "SE Asian",
    "South Asian",
    "MENA",
    "Black",
    "White",
    "Latino",
    "Indigenous",
    "Pacific Islander",
    "Mixed",
]
INCOME_MEDIAN_MULTIPLIERS = [0.75, 1.0, 1.5, 2.0]
INCOME_LEVELS = ["low", "lower-middle", "middle", "upper-middle", "high"]
EDUCATION_LEVELS = [
    "none",
    "primary",
    "lower secondary",
    "upper secondary",
    "vocational",
    "bachelor",
    "master",
    "doctorate",
]
EDUCATION_INDICATORS = {
    "primary": "SE.PRM.CUAT.ZS",
    "lower_secondary": "SE.SEC.CUAT.LO.ZS",
    "upper_secondary": "SE.SEC.CUAT.UP.ZS",
    "post_secondary": "SE.SEC.CUAT.PO.ZS",
    "bachelor": "SE.TER.CUAT.BA.ZS",
    "master": "SE.TER.CUAT.MS.ZS",
    "doctorate": "SE.TER.CUAT.DO.ZS",
}

ILO_EDUCATION_GROUPS = {
    "less than basic": {"EDU_AGGREGATE_LTB"},
    "basic": {"EDU_AGGREGATE_BAS"},
    "intermediate": {"EDU_AGGREGATE_INT"},
    "advanced": {"EDU_AGGREGATE_ADV"},
}
ILO_EDUCATION_LEVEL_TO_GROUP = {
    "none": "less than basic",
    "primary": "basic",
    "lower secondary": "basic",
    "upper secondary": "intermediate",
    "vocational": "intermediate",
    "bachelor": "advanced",
    "master": "advanced",
    "doctorate": "advanced",
}
ILO_EDUCATION_DETAILED_TO_GROUP = {
    "EDU_ISCED11_0": "less than basic",
    "EDU_ISCED11_1": "basic",
    "EDU_ISCED11_2": "basic",
    "EDU_ISCED11_3": "intermediate",
    "EDU_ISCED11_4": "intermediate",
    "EDU_ISCED11_5": "advanced",
    "EDU_ISCED11_6": "advanced",
    "EDU_ISCED11_7": "advanced",
    "EDU_ISCED11_8": "advanced",
    "EDU_ISCED97_0": "less than basic",
    "EDU_ISCED97_1": "basic",
    "EDU_ISCED97_2": "basic",
    "EDU_ISCED97_3": "intermediate",
    "EDU_ISCED97_4": "intermediate",
    "EDU_ISCED97_5": "advanced",
    "EDU_ISCED97_5A": "advanced",
    "EDU_ISCED97_5B": "advanced",
    "EDU_ISCED97_6": "advanced",
}
ILO_AGE_CODES = {
    "AGE_10YRBANDS_Y15-24": "18-24",
    "AGE_10YRBANDS_Y25-34": "25-34",
    "AGE_10YRBANDS_Y35-44": "35-44",
    "AGE_10YRBANDS_Y45-54": "45-54",
    "AGE_10YRBANDS_Y55-64": "55-64",
    "AGE_10YRBANDS_YGE65": "65-80",
}
ILO_SEX_CODES = {"SEX_M": "male", "SEX_F": "female"}

WELFARE_EDUCATION_TYPES = {
    "no education": "No education",
    "primary": "Primary (complete or incomplete)",
    "secondary": "Secondary (complete or incomplete)",
    "tertiary": "Tertiary (complete or incomplete)",
}
MIN_DIRECT_WELFARE_GROUP_SHARE_RATIO = 0.1
MAX_DIRECT_WELFARE_GROUP_SHARE_RATIO = 10.0
WELFARE_EDUCATION_LEVEL_TO_GROUP = {
    "none": "no education",
    "primary": "primary",
    "lower secondary": "secondary",
    "upper secondary": "secondary",
    "vocational": "secondary",
    "bachelor": "tertiary",
    "master": "tertiary",
    "doctorate": "tertiary",
}
WELFARE_AGE_TYPES = {
    "youth": "youth [15-24]",
    "adult": "adult [25+]",
}
ILO_AGE_STAGE_CODES = {
    "AGE_YTHADULT_Y15-24": "youth",
    "AGE_YTHADULT_YGE25": "adult",
}
ILO_MARITAL_STATUS_TO_RELATIONSHIP = {
    "MTS_DETAILS_SGLE": "single",
    "MTS_DETAILS_UNION": "in relationship",
    "MTS_DETAILS_MRD": "married",
    "MTS_DETAILS_SEP": "divorced",
    "MTS_DETAILS_WID": "widowed",
}

OCCUPATION_CODES = {
    "OCU_ISCO08_1": "managers",
    "OCU_ISCO08_2": "professionals",
    "OCU_ISCO08_3": "technicians and associate professionals",
    "OCU_ISCO08_4": "clerical support workers",
    "OCU_ISCO08_5": "service and sales workers",
    "OCU_ISCO08_7": "craft and related trades workers",
    "OCU_ISCO08_8": "plant and machine operators and assemblers",
    "OCU_ISCO08_96": "elementary and skilled agricultural workers",
}
OCCUPATIONS = ["None", *OCCUPATION_CODES.values()]
ILO_OBSERVED_OCCUPATION_TO_GROUP = {
    "OCU_ISCO08_1": "managers",
    "OCU_ISCO08_2": "professionals",
    "OCU_ISCO08_3": "technicians and associate professionals",
    "OCU_ISCO08_4": "clerical support workers",
    "OCU_ISCO08_5": "service and sales workers",
    "OCU_ISCO08_6": "elementary and skilled agricultural workers",
    "OCU_ISCO08_7": "craft and related trades workers",
    "OCU_ISCO08_8": "plant and machine operators and assemblers",
    "OCU_ISCO08_9": "elementary and skilled agricultural workers",
}

# IHME GBD 2023 prevalence query dimensions.  IDs are used instead of the
# localized display names so the import is stable across interface languages.
GBD_LOCATION_ID_TO_COUNTRY = {
    102: "United States",
    6: "China",
    163: "India",
    135: "Brazil",
    130: "Mexico",
    214: "Nigeria",
    81: "Germany",
    67: "Japan",
    4749: "United Kingdom",
    101: "Canada",
    125: "Colombia",
    80: "France",
    51: "Poland",
    68: "South Korea",
    11: "Indonesia",
    16: "Philippines",
    20: "Vietnam",
    165: "Pakistan",
    161: "Bangladesh",
    141: "Egypt",
    155: "Turkey",
    142: "Iran",
    179: "Ethiopia",
    196: "South Africa",
    180: "Kenya",
    71: "Australia",
}
GBD_SEX_ID_TO_SEX = {1: "male", 2: "female"}
GBD_AGE_ID_TO_RANGE = {
    8: (15, 19),
    9: (20, 24),
    10: (25, 29),
    11: (30, 34),
    12: (35, 39),
    13: (40, 44),
    14: (45, 49),
    15: (50, 54),
    16: (55, 59),
    17: (60, 64),
    18: (65, 69),
    19: (70, 74),
    20: (75, 79),
    30: (80, 84),
}

PHYSICAL_CAUSE_IDS = {
    "migraine": (547,),
    "tension-type headache": (548,),
    "idiopathic epilepsy": (545,),
    "Parkinson's disease": (544,),
    "Alzheimer's disease and other dementias": (543,),
    "stroke": (494,),
    "age-related and other hearing loss": (674,),
    "osteoarthritis": (628,),
    "rheumatoid arthritis": (627,),
    "low back pain": (630,),
    "asthma": (515,),
    "chronic obstructive pulmonary disease": (509,),
    "type 2 diabetes": (976,),
    "chronic kidney disease": (589,),
    "atrial fibrillation and flutter": (500,),
    "gastritis and duodenitis": (528,),
    "peptic ulcer disease": (527,),
    "dermatitis": (654,),
    "urticaria": (664,),
    "psoriasis": (655,),
    "alopecia areata": (662,),
    "HIV/AIDS": (298,),
}
MENTAL_CAUSE_IDS = {
    "schizophrenia": (559,),
    # The query exports major depressive disorder and dysthymia separately.
    "depressive disorders": (568, 569),
    "bipolar disorder": (570,),
    "anxiety disorders": (571,),
    # The query exports anorexia nervosa and bulimia nervosa separately.
    "eating disorders": (573, 574),
    "autism spectrum disorders": (575,),
    "ADHD": (578,),
    "conduct disorder": (579,),
    "idiopathic developmental intellectual disability": (582,),
    "other mental disorders": (585,),
    "alcohol use disorders": (560,),
    "opioid use disorders": (562,),
    "amphetamine use disorders": (564,),
    "cocaine use disorders": (563,),
    "cannabis use disorders": (565,),
    "other drug use disorders": (566,),
}
PHYSICAL_CONDITIONS = ["None", *PHYSICAL_CAUSE_IDS]
MENTAL_CONDITIONS = ["None", *MENTAL_CAUSE_IDS]

PHYSICAL_SYSTEMS = {
    "migraine": "neurological",
    "tension-type headache": "neurological",
    "idiopathic epilepsy": "neurological",
    "Parkinson's disease": "neurological",
    "Alzheimer's disease and other dementias": "neurological",
    "stroke": "cardiovascular_and_neurological",
    "age-related and other hearing loss": "sensory",
    "osteoarthritis": "musculoskeletal",
    "rheumatoid arthritis": "musculoskeletal",
    "low back pain": "musculoskeletal",
    "asthma": "respiratory",
    "chronic obstructive pulmonary disease": "respiratory",
    "type 2 diabetes": "metabolic",
    "chronic kidney disease": "renal",
    "atrial fibrillation and flutter": "cardiovascular",
    "gastritis and duodenitis": "digestive",
    "peptic ulcer disease": "digestive",
    "dermatitis": "skin",
    "urticaria": "skin",
    "psoriasis": "skin",
    "alopecia areata": "skin",
    "HIV/AIDS": "infectious",
}
PHYSICAL_CONDITION_STATISTICS = {
    condition: {
        "gbd_cause_ids": list(cause_ids),
        "system": PHYSICAL_SYSTEMS[condition],
        "estimate_year": 2023,
        "scope": "country_age_sex_specific",
        "span_type": "GBD_cause_or_documented_child_cause_sum",
        "source_key": "ihme_gbd_2023_prevalence_query",
    }
    for condition, cause_ids in PHYSICAL_CAUSE_IDS.items()
}
MENTAL_CONDITION_STATISTICS = {
    condition: {
        "gbd_cause_ids": list(cause_ids),
        "family": (
            "alcohol_use_disorders"
            if condition == "alcohol use disorders"
            else (
                "drug_use_disorders"
                if condition.endswith("drug use disorders")
                or condition
                in {
                    "opioid use disorders",
                    "amphetamine use disorders",
                    "cocaine use disorders",
                    "cannabis use disorders",
                }
                else "mental_disorders"
            )
        ),
        "estimate_year": 2023,
        "scope": "country_age_sex_specific",
        "span_type": "GBD_cause_or_documented_child_cause_sum",
        "source_key": "ihme_gbd_2023_prevalence_query",
    }
    for condition, cause_ids in MENTAL_CAUSE_IDS.items()
}

# Non-overlapping IMS origin subregions, mapped to the generator's broad
# ethnicity taxonomy.  This is used only for countries without a compatible
# official population-group table.
IMS_ORIGIN_TO_ETHNICITY = {
    910: "Black",  # Eastern Africa
    911: "Black",  # Middle Africa
    913: "Black",  # Southern Africa
    914: "Black",  # Western Africa
    912: "MENA",  # Northern Africa
    5500: "MENA",  # Central Asia
    922: "MENA",  # Western Asia
    5501: "South Asian",  # Southern Asia
    906: "East Asian",  # Eastern Asia
    920: "SE Asian",  # South-Eastern Asia
    915: "Latino",  # Caribbean
    916: "Latino",  # Central America
    931: "Latino",  # South America
    923: "White",  # Eastern Europe
    924: "White",  # Northern Europe
    925: "White",  # Southern Europe
    926: "White",  # Western Europe
    905: "White",  # Northern America
    927: "White",  # Australia and New Zealand
    928: "Pacific Islander",  # Melanesia
    954: "Pacific Islander",  # Micronesia
    957: "Pacific Islander",  # Polynesia
    2003: "Mixed",  # Other/unknown origin
}
BASE_ETHNICITY = {
    "Colombia": "Latino",
    "France": "White",
    "Germany": "White",
    "Poland": "White",
    "South Korea": "East Asian",
    "Indonesia": "SE Asian",
    "Philippines": "SE Asian",
    "Vietnam": "SE Asian",
    "India": "South Asian",
    "Pakistan": "South Asian",
    "Bangladesh": "South Asian",
    "Egypt": "MENA",
    "Turkey": "MENA",
    "Iran": "MENA",
    "Japan": "East Asian",
    "Nigeria": "Black",
    "Ethiopia": "Black",
    "South Africa": "Black",
    "Kenya": "Black",
    "Australia": "White",
}


def download(
    url: str, path: Path, *, headers: dict[str, str] | None = None
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and path.stat().st_size > 200:
        return path
    request_headers = {
        "User-Agent": "Mozilla/5.0",
        "Referer": "https://www.un.org/development/desa/pd/content/international-migrant-stock",
    }
    request_headers.update(headers or {})
    request = urllib.request.Request(url, headers=request_headers)
    with urllib.request.urlopen(request, timeout=240) as response:
        path.write_bytes(response.read())
    return path


def read_zipped_csv(path: Path) -> list[dict[str, str]]:
    with zipfile.ZipFile(path) as archive:
        name = next(name for name in archive.namelist() if name.endswith(".csv"))
        text = io.TextIOWrapper(archive.open(name), encoding="utf-8-sig")
        return list(csv.DictReader(text))


def overlap_fraction(low: int, high: int, target_low: int, target_high: int) -> float:
    overlap = max(0, min(high, target_high) - max(low, target_low) + 1)
    return overlap / (high - low + 1)


def parse_age_range(label: str) -> tuple[int, int] | None:
    match = re.fullmatch(r"\s*(\d+)\s*-\s*(\d+)\s*", label)
    if match:
        return int(match.group(1)), int(match.group(2))
    match = re.fullmatch(r"\s*(\d+)\s*\+\s*", label)
    if match:
        return int(match.group(1)), 120
    return None


def normalize(values: dict[str, float]) -> dict[str, float]:
    total = sum(max(0.0, value) for value in values.values())
    if total <= 0:
        raise ValueError("Cannot normalize an empty distribution")
    return {key: max(0.0, value) / total for key, value in values.items()}


def rake_matrix(
    prior: dict[tuple[str, str], float],
    row_targets: dict[str, float],
    column_targets: dict[str, float],
    *,
    tolerance: float = 1e-10,
    max_iterations: int = 2_000,
) -> dict[tuple[str, str], float]:
    """Rake a non-negative two-dimensional prior to exact row/column margins."""

    rows = list(row_targets)
    columns = list(column_targets)
    row_total = sum(row_targets.values())
    column_total = sum(column_targets.values())
    if not math.isclose(row_total, column_total, rel_tol=0, abs_tol=1e-9):
        raise ValueError(
            f"Raking margins disagree: rows={row_total}, columns={column_total}"
        )
    scale = max(row_total, 1.0)
    epsilon = scale * 1e-12
    values = {
        (row, column): (
            max(0.0, float(prior.get((row, column), 0.0))) + epsilon
            if row_targets[row] > 0 and column_targets[column] > 0
            else 0.0
        )
        for row in rows
        for column in columns
    }

    for _ in range(max_iterations):
        for row in rows:
            current = sum(values[(row, column)] for column in columns)
            target = row_targets[row]
            factor = target / current if current > 0 else 0.0
            for column in columns:
                values[(row, column)] *= factor
        for column in columns:
            current = sum(values[(row, column)] for row in rows)
            target = column_targets[column]
            factor = target / current if current > 0 else 0.0
            for row in rows:
                values[(row, column)] *= factor
        maximum_error = max(
            [
                abs(
                    sum(values[(row, column)] for column in columns)
                    - row_targets[row]
                )
                for row in rows
            ]
            + [
                abs(
                    sum(values[(row, column)] for row in rows)
                    - column_targets[column]
                )
                for column in columns
            ]
        )
        if maximum_error <= tolerance:
            return values
    raise RuntimeError(
        f"Raking did not converge after {max_iterations} iterations; "
        f"max error={maximum_error}"
    )


def calibrate_probabilities_to_mean(
    probabilities: dict[str, float],
    weights: dict[str, float],
    target_mean: float,
) -> dict[str, float]:
    """Apply one logit intercept shift while preserving probability ordering."""

    target_mean = min(1.0, max(0.0, target_mean))
    normalized_weights = normalize(weights)
    if target_mean <= 0:
        return {key: 0.0 for key in probabilities}
    if target_mean >= 1:
        return {key: 1.0 for key in probabilities}

    logits = {}
    for key, probability in probabilities.items():
        probability = min(1.0 - 1e-9, max(1e-9, float(probability)))
        logits[key] = math.log(probability / (1.0 - probability))

    low, high = -40.0, 40.0
    calibrated: dict[str, float] = {}
    for _ in range(200):
        shift = (low + high) / 2.0
        calibrated = {
            key: 1.0 / (1.0 + math.exp(-(logit + shift)))
            for key, logit in logits.items()
        }
        mean = sum(
            normalized_weights[key] * calibrated[key]
            for key in probabilities
        )
        if mean < target_mean:
            low = shift
        else:
            high = shift
    return calibrated


def grouped_education_margins(
    education_targets: dict[str, float],
    country: str,
    mapping: dict[str, str],
) -> dict[str, float]:
    groups = list(dict.fromkeys(mapping.values()))
    values = {group: 0.0 for group in groups}
    for education_level, group in mapping.items():
        values[group] += education_targets[
            f"{country} | {education_level}"
        ]
    return values


def _ilo_candidate_groups(frame: Any) -> list[tuple[int, float, str, Any]]:
    candidates = []
    for (time_value, source), group in frame.groupby(["time", "source"]):
        year = int(time_value)
        if year > 2025:
            continue
        best = float(group["best_source"].max())
        candidates.append((year, best, str(source), group))
    return sorted(candidates, key=lambda item: (item[0], item[1]), reverse=True)


def _education_code_mapping(mode: str) -> dict[str, str]:
    if mode == "aggregate":
        return {
            code: group
            for group, codes in ILO_EDUCATION_GROUPS.items()
            for code in codes
        }
    prefix = "EDU_ISCED11_" if mode == "isced11" else "EDU_ISCED97_"
    return {
        code: group
        for code, group in ILO_EDUCATION_DETAILED_TO_GROUP.items()
        if code.startswith(prefix)
    }


def select_ilo_age_education_panel(
    frame: Any,
    iso3: str,
) -> tuple[dict[tuple[str, str, str], float], dict[str, Any]] | None:
    """Select the latest complete sex × 10-year-age × education panel."""

    subset = frame[
        (frame["ref_area"] == iso3)
        & (frame["sex"].isin(ILO_SEX_CODES))
        & (frame["classif1"].isin(ILO_AGE_CODES))
    ]
    required = {
        (age_group, sex, education_group)
        for age_group in AGE_GROUPS
        for sex in ILO_SEX_CODES.values()
        for education_group in ILO_EDUCATION_GROUPS
    }
    for year, best, source, candidate in _ilo_candidate_groups(subset):
        for mode in ("aggregate", "isced11", "isced97"):
            code_mapping = _education_code_mapping(mode)
            rows = candidate[candidate["classif2"].isin(code_mapping)]
            values: dict[tuple[str, str, str], float] = defaultdict(float)
            present: set[tuple[str, str, str]] = set()
            for row in rows.itertuples(index=False):
                key = (
                    ILO_AGE_CODES[str(row.classif1)],
                    ILO_SEX_CODES[str(row.sex)],
                    code_mapping[str(row.classif2)],
                )
                values[key] += max(0.0, float(row.obs_value))
                present.add(key)
            if present == required:
                return dict(values), {
                    "year": year,
                    "source": source,
                    "best_source": bool(best),
                    "education_classification": mode,
                }
    return None


def select_ilo_age_rate_panel(
    frame: Any,
    iso3: str,
) -> tuple[dict[tuple[str, str], float], dict[str, Any]] | None:
    """Select the latest complete sex × 10-year-age employment-rate panel."""

    subset = frame[
        (frame["ref_area"] == iso3)
        & (frame["sex"].isin(ILO_SEX_CODES))
        & (frame["classif1"].isin(ILO_AGE_CODES))
    ]
    required = {
        (age_group, sex)
        for age_group in AGE_GROUPS
        for sex in ILO_SEX_CODES.values()
    }
    for year, best, source, candidate in _ilo_candidate_groups(subset):
        values: dict[tuple[str, str], float] = {}
        for row in candidate.itertuples(index=False):
            key = (
                ILO_AGE_CODES[str(row.classif1)],
                ILO_SEX_CODES[str(row.sex)],
            )
            values[key] = min(1.0, max(0.0, float(row.obs_value) / 100.0))
        if set(values) == required:
            return values, {
                "year": year,
                "source": source,
                "best_source": bool(best),
            }
    return None


def load_ilo_rds(path: Path) -> Any:
    """Read an ILOSTAT bulk RDS without requiring an R installation."""

    try:
        import rdata
        from rdata.conversion import DEFAULT_CLASS_MAP
    except ImportError as exc:
        raise RuntimeError(
            "The official ILOSTAT bulk files require the 'rdata' package. "
            "Install project requirements and rerun."
        ) from exc

    def forgiving_factor(values: Any, attrs: dict[str, Any]) -> list[Any]:
        raw_levels = attrs.get("levels")
        levels = [] if raw_levels is None else list(raw_levels)
        output = []
        for code in values:
            try:
                index = int(code) - 1
            except (TypeError, ValueError):
                output.append(None)
                continue
            output.append(levels[index] if 0 <= index < len(levels) else None)
        return output

    constructors = dict(DEFAULT_CLASS_MAP)
    constructors.update(factor=forgiving_factor, ordered=forgiving_factor)
    return rdata.read_rds(path, constructor_dict=constructors)


def constraint(data: dict[str, Any], name: str) -> dict[str, Any]:
    return next(item for item in data["constraints"] if item["name"] == name)


def load_wpp(path: Path) -> dict[str, Any]:
    age_sex: dict[tuple[str, int, str, str], float] = defaultdict(float)
    age5_sex: dict[tuple[str, int, int, int, str], float] = defaultdict(float)
    total: dict[tuple[str, int], float] = defaultdict(float)
    total_sex: dict[tuple[str, int, str], float] = defaultdict(float)
    adult: dict[tuple[str, int], float] = defaultdict(float)
    supported = {info["iso3"]: country for country, info in COUNTRY_CODES.items()}
    with gzip.open(path, "rt", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            if row["ISO3_code"] not in supported or int(row["Time"]) not in {2020, 2025}:
                continue
            country = supported[row["ISO3_code"]]
            year = int(row["Time"])
            low = int(row["AgeGrpStart"])
            span = int(row["AgeGrpSpan"])
            high = low + span - 1
            male = float(row["PopMale"]) * 1000
            female = float(row["PopFemale"]) * 1000
            age5_sex[(country, year, low, high, "male")] += male
            age5_sex[(country, year, low, high, "female")] += female
            total[(country, year)] += male + female
            total_sex[(country, year, "male")] += male
            total_sex[(country, year, "female")] += female
            adult_fraction = overlap_fraction(low, high, 18, 80)
            adult[(country, year)] += (male + female) * adult_fraction
            if year == 2025:
                for age_group, (target_low, target_high) in AGE_GROUPS.items():
                    fraction = overlap_fraction(low, high, target_low, target_high)
                    if fraction:
                        age_sex[(country, year, age_group, "male")] += male * fraction
                        age_sex[(country, year, age_group, "female")] += female * fraction
    return {
        "age_sex": age_sex,
        "age5_sex": age5_sex,
        "total": total,
        "total_sex": total_sex,
        "adult": adult,
    }


def build_population_targets(
    wpp: dict[str, Any],
) -> tuple[dict[str, float], dict[str, float]]:
    country_population = normalize(
        {
            country: wpp["adult"][(country, 2025)]
            for country in COUNTRY_CODES
        }
    )
    age_sex_population = {
        f"{age_group} | {sex}": sum(
            wpp["age_sex"][(country, 2025, age_group, sex)]
            for country in COUNTRY_CODES
        )
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    }
    return country_population, normalize(age_sex_population)


def build_country_age_sex_targets(wpp: dict[str, Any]) -> dict[str, float]:
    population = {
        f"{country} | {age_group} | {sex}": wpp["age_sex"][
            (country, 2025, age_group, sex)
        ]
        for country in COUNTRY_CODES
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    }
    return normalize(population)


def city_size_class(population: float) -> str:
    if population < 100_000:
        return "50k-100k"
    if population < 250_000:
        return "100k-250k"
    if population < 500_000:
        return "250k-500k"
    if population < 1_000_000:
        return "500k-1m"
    if population < 5_000_000:
        return "1m-5m"
    if population < 10_000_000:
        return "5m-10m"
    return "10m+"


def load_wup_city_catalog(
    path: Path,
    country_mass: dict[str, float],
) -> tuple[dict[str, Any], dict[str, float]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Data"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    index = {str(value): position for position, value in enumerate(header)}
    iso_to_country = {
        info["iso3"]: country for country, info in COUNTRY_CODES.items()
    }
    catalog: dict[str, list[dict[str, Any]]] = {
        country: [] for country in COUNTRY_CODES
    }
    for row in rows:
        iso3 = row[index["ISO3_Code"]]
        country = iso_to_country.get(str(iso3))
        population_thousands = row[index["2025"]]
        city_name = row[index["City_Name"]]
        if (
            country is None
            or city_name is None
            or population_thousands is None
            or float(population_thousands) < 50.0
        ):
            continue
        population = round(float(population_thousands) * 1000)
        catalog[country].append(
            {
                "name": str(city_name).strip(),
                "population_2025": population,
                "size_class": city_size_class(population),
            }
        )

    targets: dict[str, float] = {}
    for country, cities in catalog.items():
        if not cities:
            raise RuntimeError(f"WUP returned no 50,000+ cities for {country}.")
        cities.sort(key=lambda item: (-item["population_2025"], item["name"]))
        by_size: dict[str, float] = defaultdict(float)
        for city in cities:
            by_size[city["size_class"]] += city["population_2025"]
        conditional = normalize(by_size)
        for size_class, share in conditional.items():
            targets[f"{country} | {size_class}"] = country_mass[country] * share

    output = {
        "source": WUP_CITY_URL,
        "year": 2025,
        "minimum_population": 50_000,
        "countries": catalog,
    }
    return output, targets


def fetch_wdi_education(cache: Path) -> dict[str, Any]:
    cache_file = cache / "world_bank_education_attainment.json"
    if cache_file.exists():
        cached = json.loads(cache_file.read_text(encoding="utf-8"))
    else:
        cached = {}
    iso_to_country = {
        info["iso3"]: country for country, info in COUNTRY_CODES.items()
    }
    country_path = ";".join(iso_to_country)
    for label, indicator in EDUCATION_INDICATORS.items():
        if all(label in cached.get(country, {}) for country in COUNTRY_CODES):
            continue
        url = WORLD_BANK_API.format(
            countries=country_path,
            indicator=indicator,
        )
        with urllib.request.urlopen(url, timeout=240) as response:
            payload = json.load(response)
        rows = payload[1] if isinstance(payload, list) and len(payload) > 1 else []
        seen: set[str] = set()
        for row in rows:
            iso3 = str(row.get("countryiso3code", ""))
            country = iso_to_country.get(iso3)
            if (
                country is None
                or country in seen
                or row.get("value") is None
            ):
                continue
            cached.setdefault(country, {})[label] = {
                "year": int(row["date"]),
                "value": float(row["value"]) / 100.0,
                "indicator": indicator,
            }
            seen.add(country)
    cache_file.write_text(
        json.dumps(cached, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return cached


def build_education_targets(
    country_mass: dict[str, float],
    observations: dict[str, Any],
) -> tuple[dict[str, float], dict[str, Any]]:
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for country, mass in country_mass.items():
        values = {
            label: min(1.0, max(0.0, float(item["value"])))
            for label, item in observations[country].items()
        }
        primary = values["primary"]
        lower = min(primary, values["lower_secondary"])
        upper = min(lower, values["upper_secondary"])
        post = min(upper, values.get("post_secondary", 0.0))
        bachelor = min(upper - post, values["bachelor"])
        master = min(bachelor, values["master"])
        doctorate = min(master, values.get("doctorate", 0.0))
        shares = {
            "none": 1.0 - primary,
            "primary": primary - lower,
            "lower secondary": lower - upper,
            "upper secondary": upper - post - bachelor,
            "vocational": post,
            "bachelor": bachelor - master,
            "master": master - doctorate,
            "doctorate": doctorate,
        }
        shares = normalize(shares)
        for level in EDUCATION_LEVELS:
            targets[f"{country} | {level}"] = mass * shares[level]
        audit[country] = {
            "indicator_observations": observations[country],
            "education_level_shares": shares,
        }
    return targets, audit


def load_pew_religion(
    path: Path,
    country_mass: dict[str, float],
) -> tuple[dict[str, float], dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        name = next(
            item
            for item in archive.namelist()
            if item.endswith("(percentages).csv")
        )
        text = io.TextIOWrapper(archive.open(name), encoding="utf-8-sig")
        rows = list(csv.DictReader(text))
    m49_to_country = {
        info["m49"]: country for country, info in COUNTRY_CODES.items()
    }
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for row in rows:
        if row.get("Year") != "2020":
            continue
        try:
            country = m49_to_country[int(row["Countrycode"])]
        except (KeyError, TypeError, ValueError):
            continue
        shares = normalize(
            {
                "Christian": float(row["Christians"]),
                "Muslim": float(row["Muslims"]),
                "Unaffiliated": float(row["Religiously_unaffiliated"]),
                "Buddhist": float(row["Buddhists"]),
                "Hindu": float(row["Hindus"]),
                "Jewish": float(row["Jews"]),
                "Folk": float(row["Other_religions"]),
            }
        )
        for religion, share in shares.items():
            targets[f"{country} | {religion}"] = (
                country_mass[country] * share
            )
        audit[country] = {
            "source_country": row["Country"],
            "year": 2020,
            "religious_belief_shares": shares,
        }
    missing = sorted(set(COUNTRY_CODES) - set(audit))
    if missing:
        raise RuntimeError(f"Pew religion data missing countries: {missing}")
    return targets, audit


def load_ims(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Table 1"]
    code_to_country = {info["m49"]: country for country, info in COUNTRY_CODES.items()}
    destinations: dict[str, dict[str, Any]] = {}
    wanted_origins = set(code_to_country) | set(IMS_ORIGIN_TO_ETHNICITY) | {900}
    for row in sheet.iter_rows(min_row=12, values_only=True):
        destination_code = row[3]
        origin_code = row[6]
        if destination_code not in code_to_country or origin_code not in wanted_origins:
            continue
        country = code_to_country[destination_code]
        entry = destinations.setdefault(
            country,
            {
                "data_type": str(row[4] or ""),
                "origins": {},
                "origins_by_sex": {"male": {}, "female": {}},
            },
        )
        entry["origins"][int(origin_code)] = float(row[13] or 0)
        entry["origins_by_sex"]["male"][int(origin_code)] = float(
            row[20] or 0
        )
        entry["origins_by_sex"]["female"][int(origin_code)] = float(
            row[27] or 0
        )
    return destinations


def build_migration_targets(
    country_mass: dict[str, float],
    wpp: dict[str, Any],
    ims: dict[str, Any],
) -> tuple[dict[str, float], dict[str, Any]]:
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for destination, mass in country_mass.items():
        origins = ims[destination]["origins"]
        population = wpp["total"][(destination, 2020)]
        migrant_total = origins[900]
        native = max(0.0, population - migrant_total)
        supported = {
            origin: origins.get(info["m49"], 0.0)
            for origin, info in COUNTRY_CODES.items()
            if origin != destination
        }
        denominator = native + sum(supported.values())
        row = {origin: value / denominator for origin, value in supported.items()}
        row[destination] = native / denominator
        for origin in COUNTRY_CODES:
            targets[f"{destination} | {origin}"] = mass * row.get(origin, 0.0)
        basis = ims[destination]["data_type"]
        audit[destination] = {
            "ims_data_type": basis,
            "direct_for_birth_country": "B" in basis,
            "direct_for_citizenship_country": "C" in basis,
            "wpp_2020_population": round(population),
            "ims_2020_migrant_total": round(migrant_total),
            "supported_foreign_origin_total": round(sum(supported.values())),
            "conditioning": (
                f"native plus foreign origins in the supported "
                f"{len(COUNTRY_CODES)}-country universe"
            ),
        }
    return targets, audit


def build_migration_status_by_sex_targets(
    wpp: dict[str, Any],
    ims: dict[str, Any],
) -> tuple[dict[str, float], dict[str, Any]]:
    """Build current-country × sex × local/foreign status from IMS Table 1."""

    total_adult_population = sum(
        wpp["age_sex"][(country, 2025, age_group, sex)]
        for country in COUNTRY_CODES
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    )
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for country in COUNTRY_CODES:
        audit[country] = {}
        for sex in ("female", "male"):
            origins = ims[country]["origins_by_sex"][sex]
            migrant_total = origins[900]
            native = max(
                0.0,
                wpp["total_sex"][(country, 2020, sex)] - migrant_total,
            )
            supported_foreign = sum(
                origins.get(info["m49"], 0.0)
                for origin, info in COUNTRY_CODES.items()
                if origin != country
            )
            conditional = normalize(
                {"local": native, "foreign": supported_foreign}
            )
            sex_mass = sum(
                wpp["age_sex"][(country, 2025, age_group, sex)]
                for age_group in AGE_GROUPS
            ) / total_adult_population
            for status, share in conditional.items():
                targets[f"{country} | {sex} | {status}"] = (
                    sex_mass * share
                )
            audit[country][sex] = {
                "wpp_2020_same_sex_population": round(
                    wpp["total_sex"][(country, 2020, sex)]
                ),
                "ims_2020_same_sex_migrant_total": round(migrant_total),
                "supported_foreign_origin_total": round(supported_foreign),
                "conditional_status_shares": conditional,
            }
    return targets, audit


def load_marital_rows(cache: Path) -> tuple[dict[str, Any], dict[str, int]]:
    country_rates: dict[str, Any] = {}
    years: dict[str, int] = {}
    for country, info in COUNTRY_CODES.items():
        url = UNDATA_URL.format(table=23, country_code=info["m49"])
        try:
            rows = read_zipped_csv(
                download(url, cache / f"undata_23_{info['m49']}.zip")
            )
        except Exception:
            continue
        rows = [
            row
            for row in rows
            if row.get("Year", "").isdigit()
            and row.get("Area") == "Total"
            and row.get("Sex") in SEX_MAP
        ]
        if not rows:
            continue
        latest = max(int(row["Year"]) for row in rows)
        years[country] = latest
        rows = [row for row in rows if int(row["Year"]) == latest]
        rates: dict[tuple[str, str], dict[str, float]] = {}
        for age_group, (target_low, target_high) in AGE_GROUPS.items():
            for source_sex, sex in SEX_MAP.items():
                counts = {status: 0.0 for status in set(RELATIONSHIP_MAP.values())}
                for row in rows:
                    if row["Sex"] != source_sex or row["Marital status"] not in RELATIONSHIP_MAP:
                        continue
                    age_range = parse_age_range(row["Age"])
                    if not age_range:
                        continue
                    fraction = overlap_fraction(*age_range, target_low, target_high)
                    if fraction:
                        counts[RELATIONSHIP_MAP[row["Marital status"]]] += (
                            float(row["Value"]) * fraction
                        )
                rates[(age_group, sex)] = normalize(counts)
        country_rates[country] = rates
    return country_rates, years


def build_relationship_targets(
    age_sex_mass: dict[str, float],
    wpp: dict[str, Any],
    country_rates: dict[str, Any],
) -> dict[str, float]:
    targets: dict[str, float] = {}
    statuses = ["single", "in relationship", "married", "divorced", "widowed"]
    for age_group in AGE_GROUPS:
        for sex in ("female", "male"):
            weighted = {status: 0.0 for status in statuses}
            denominator = 0.0
            for country in COUNTRY_CODES:
                if country not in country_rates:
                    continue
                population = wpp["age_sex"][(country, 2025, age_group, sex)]
                denominator += population
                for status in statuses:
                    weighted[status] += (
                        population * country_rates[country][(age_group, sex)].get(status, 0.0)
                    )
            conditional = normalize(weighted)
            row_mass = age_sex_mass[f"{age_group} | {sex}"]
            for status in statuses:
                targets[f"{age_group} | {sex} | {status}"] = (
                    row_mass * conditional[status]
                )
    return targets


def build_country_relationship_targets(
    wpp: dict[str, Any],
    country_rates: dict[str, Any],
) -> dict[str, float]:
    """Retain the official country dimension instead of pooling it away."""

    statuses = ["single", "in relationship", "married", "divorced", "widowed"]
    total_population = sum(
        wpp["age_sex"][(country, 2025, age_group, sex)]
        for country in COUNTRY_CODES
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    )
    targets: dict[str, float] = {}
    for country in COUNTRY_CODES:
        if country not in country_rates:
            raise RuntimeError(
                f"UNData marital-status rows are missing for {country}."
            )
        for age_group in AGE_GROUPS:
            for sex in ("female", "male"):
                row_mass = (
                    wpp["age_sex"][(country, 2025, age_group, sex)]
                    / total_population
                )
                conditional = country_rates[country][(age_group, sex)]
                for status in statuses:
                    targets[
                        f"{country} | {age_group} | {sex} | {status}"
                    ] = row_mass * conditional.get(status, 0.0)
    return targets


def latest_national_pip_row(
    rows: list[dict[str, Any]], *, require_median: bool = False
) -> dict[str, Any]:
    eligible = [
        row
        for row in rows
        if row.get("reporting_level") == "national"
        and row.get("headcount") is not None
        and (not require_median or row.get("median") is not None)
    ]
    if not eligible:
        raise RuntimeError("PIP returned no eligible national observations.")
    observed = [row for row in eligible if not bool(row.get("is_interpolated"))]
    return max(
        observed or eligible,
        key=lambda row: (
            int(row["reporting_year"]),
            int(row.get("survey_year") or 0),
        ),
    )


def matching_pip_row(
    rows: list[dict[str, Any]], reference: dict[str, Any]
) -> dict[str, Any]:
    identity_fields = (
        "reporting_year",
        "survey_year",
        "reporting_level",
        "welfare_type",
        "is_interpolated",
    )
    matches = [
        row
        for row in rows
        if row.get("headcount") is not None
        and all(row.get(field) == reference.get(field) for field in identity_fields)
    ]
    if not matches:
        raise RuntimeError(
            "PIP did not return the selected distribution at the requested threshold."
        )
    return matches[0]


def fetch_pip_rows(cache: Path) -> dict[str, Any]:
    cache_file = cache / "pip_2021_ppp_relative_median_headcounts.json"
    if cache_file.exists():
        output: dict[str, dict[str, Any]] = json.loads(
            cache_file.read_text(encoding="utf-8")
        )
    else:
        output = {}

    def request_rows(iso3: str, year: str | int, line: float) -> list[dict[str, Any]]:
        url = PIP_URL.format(country=iso3, year=year, line=f"{line:.8f}")
        last_error: Exception | None = None
        for _ in range(4):
            try:
                request = urllib.request.Request(
                    url, headers={"User-Agent": "Mozilla/5.0"}
                )
                with urllib.request.urlopen(request, timeout=240) as response:
                    return json.load(response)
            except Exception as exc:
                last_error = exc
        raise RuntimeError(f"PIP request failed for {iso3}: {last_error}")

    def fetch_country(country: str, iso3: str) -> tuple[str, dict[str, Any]]:
        reference = latest_national_pip_row(
            request_rows(iso3, "all", 1.0), require_median=True
        )
        median = float(reference["median"])
        thresholds = {
            str(multiplier): median * multiplier
            for multiplier in INCOME_MEDIAN_MULTIPLIERS
        }
        headcounts = {}
        for multiplier, threshold in thresholds.items():
            rows = request_rows(iso3, int(reference["reporting_year"]), threshold)
            headcounts[multiplier] = float(
                matching_pip_row(rows, reference)["headcount"]
            )
        return country, {
            "reporting_year": int(reference["reporting_year"]),
            "survey_year": reference.get("survey_year"),
            "is_interpolated": bool(reference.get("is_interpolated")),
            "welfare_type": reference["welfare_type"],
            "median_2021_ppp_usd_per_day": median,
            "thresholds_2021_ppp_usd_per_day": thresholds,
            "headcounts": headcounts,
        }

    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = {}
        for country, info in COUNTRY_CODES.items():
            if country in output:
                continue
            future = executor.submit(fetch_country, country, info["iso3"])
            futures[future] = country
        for future in as_completed(futures):
            country, row = future.result()
            output[country] = row
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(
        json.dumps(output, indent=2, sort_keys=True), encoding="utf-8"
    )
    return output


def build_income_targets(
    country_mass: dict[str, float], pip_rows: dict[str, Any]
) -> tuple[dict[str, float], dict[str, Any]]:
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for country, mass in country_mass.items():
        row = pip_rows[country]
        headcounts = [
            float(row["headcounts"][str(multiplier)])
            for multiplier in INCOME_MEDIAN_MULTIPLIERS
        ]
        shares = [
            headcounts[0],
            headcounts[1] - headcounts[0],
            headcounts[2] - headcounts[1],
            headcounts[3] - headcounts[2],
            1.0 - headcounts[3],
        ]
        shares = list(normalize({str(i): value for i, value in enumerate(shares)}).values())
        for label, share in zip(INCOME_LEVELS, shares):
            targets[f"{country} | {label}"] = mass * share
        audit[country] = {
            "reporting_year": row["reporting_year"],
            "survey_year": row.get("survey_year"),
            "is_interpolated": row.get("is_interpolated", False),
            "welfare_type": row["welfare_type"],
            "median_2021_ppp_usd_per_day": row[
                "median_2021_ppp_usd_per_day"
            ],
            "thresholds_2021_ppp_usd_per_day": row[
                "thresholds_2021_ppp_usd_per_day"
            ],
            "headcounts_by_median_multiplier": row["headcounts"],
            "income_level_shares": {
                label: share for label, share in zip(INCOME_LEVELS, shares)
            },
        }
    return targets, audit


def load_world_bank_welfare_profile(path: Path) -> Any:
    try:
        import pandas
    except ImportError as exc:
        raise RuntimeError(
            "The World Bank education-welfare Stata file requires pandas. "
            "Install project requirements and rerun."
        ) from exc
    with zipfile.ZipFile(path) as archive:
        member = next(
            (
                name
                for name in archive.namelist()
                if name.endswith(
                    "global_profile_lineup_2022_country_level.dta"
                )
            ),
            None,
        )
        if member is None:
            raise RuntimeError(
                "World Bank reproducibility package is missing the welfare "
                "profile Stata file."
            )
        return pandas.read_stata(
            io.BytesIO(archive.read(member)),
            convert_categoricals=False,
        )


def build_conditioned_income_targets(
    *,
    country_mass: dict[str, float],
    conditioned_margins: dict[str, float],
    income_targets: dict[str, float],
    welfare_rows: Any,
    group_source_types: dict[str, str],
    dimension_name: str,
) -> tuple[dict[str, float], dict[str, Any]]:
    """Transfer a World Bank conditioned welfare profile to relative bands.

    The source identifies group-specific shares below $2.15 and $6.85. PIP
    identifies the mass in this project's five median-relative bands. Within
    each observed welfare-rank interval, maximum entropy is the
    least-informative allocation that preserves both sets of margins.
    """

    source_types = set(group_source_types.values())
    frame = welfare_rows[welfare_rows["type"].isin(source_types)].copy()
    direct: dict[tuple[str, str], dict[str, Any]] = {}
    for row in frame.itertuples(index=False):
        direct[(str(row.country_code), str(row.type))] = {
            "region": str(row.region),
            "hc215": min(1.0, max(0.0, float(row.hc215))),
            "hc685": min(1.0, max(0.0, float(row.hc685))),
            "pop": max(0.0, float(row.pop)),
        }

    region_rates: dict[tuple[str, str], dict[str, float]] = {}
    for region in sorted(str(value) for value in frame["region"].unique()):
        for source_type in source_types:
            rows = [
                row
                for (iso3, row_type), row in direct.items()
                if row_type == source_type and row["region"] == region
            ]
            population = sum(row["pop"] for row in rows)
            if population <= 0:
                continue
            region_rates[(region, source_type)] = {
                "hc215": sum(
                    row["pop"] * row["hc215"] for row in rows
                )
                / population,
                "hc685": sum(
                    row["pop"] * row["hc685"] for row in rows
                )
                / population,
            }

    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for country, info in COUNTRY_CODES.items():
        iso3 = info["iso3"]
        country_rows = [
            row for (code, _), row in direct.items() if code == iso3
        ]
        source_country_population = sum(row["pop"] for row in country_rows)
        region = country_rows[0]["region"] if country_rows else "OHI"
        group_margins = {
            group: conditioned_margins[f"{country} | {group}"]
            for group in group_source_types
        }
        group_shares = {
            group: value / country_mass[country]
            for group, value in group_margins.items()
        }
        direct_share_ratios = {}
        for group, source_type in group_source_types.items():
            source_row = direct.get((iso3, source_type))
            source_share = (
                source_row["pop"] / source_country_population
                if source_row is not None and source_country_population > 0
                else None
            )
            margin_share = group_shares[group]
            direct_share_ratios[group] = (
                source_share / margin_share
                if source_share is not None and margin_share > 1e-12
                else None
            )
        country_dimension_reliable = all(
            direct.get((iso3, source_type)) is not None
            and (
                direct_share_ratios[group] is None
                or (
                    MIN_DIRECT_WELFARE_GROUP_SHARE_RATIO
                    <= direct_share_ratios[group]
                    <= MAX_DIRECT_WELFARE_GROUP_SHARE_RATIO
                )
            )
            for group, source_type in group_source_types.items()
        )
        conditional: dict[str, dict[str, float]] = {}
        provenance: dict[str, str] = {}
        source_population_shares: dict[str, float | None] = {}
        source_to_margin_share_ratios: dict[str, float | None] = {}
        for group, source_type in group_source_types.items():
            row = direct.get((iso3, source_type))
            source_share = (
                row["pop"] / source_country_population
                if row is not None and source_country_population > 0
                else None
            )
            margin_share = group_shares[group]
            share_ratio = direct_share_ratios[group]
            source_population_shares[group] = source_share
            source_to_margin_share_ratios[group] = share_ratio
            direct_row_reliable = (
                row is not None
                and country_dimension_reliable
                and (
                    share_ratio is None
                    or (
                        MIN_DIRECT_WELFARE_GROUP_SHARE_RATIO
                        <= share_ratio
                        <= MAX_DIRECT_WELFARE_GROUP_SHARE_RATIO
                    )
                )
            )
            if direct_row_reliable:
                conditional[group] = {
                    "hc215": row["hc215"],
                    "hc685": max(row["hc215"], row["hc685"]),
                }
                provenance[group] = "country_direct"
            else:
                proxy = region_rates.get((region, source_type))
                if proxy is None:
                    proxy = region_rates[("OHI", source_type)]
                conditional[group] = {
                    "hc215": proxy["hc215"],
                    "hc685": max(proxy["hc215"], proxy["hc685"]),
                }
                reason = (
                    "country_dimension_population_share_mismatch"
                    if row is not None and not country_dimension_reliable
                    else "direct_group_population_share_mismatch"
                    if row is not None
                    else "country_group_missing"
                )
                provenance[group] = (
                    f"world_bank_region_proxy:{region}:{reason}"
                )

        q215 = sum(
            group_shares[group] * conditional[group]["hc215"]
            for group in group_source_types
        )
        q685 = sum(
            group_shares[group] * conditional[group]["hc685"]
            for group in group_source_types
        )
        q685 = max(q215, q685)
        welfare_zones = [
            (0.0, q215, "below_2.15"),
            (q215, q685, "2.15_to_6.85"),
            (q685, 1.0, "above_6.85"),
        ]

        income_shares = {
            level: income_targets[f"{country} | {level}"]
            / country_mass[country]
            for level in INCOME_LEVELS
        }
        income_intervals = []
        lower = 0.0
        for level in INCOME_LEVELS:
            upper = lower + income_shares[level]
            income_intervals.append((lower, upper, level))
            lower = upper
        income_intervals[-1] = (
            income_intervals[-1][0],
            1.0,
            income_intervals[-1][2],
        )

        country_joint = {
            (group, level): 0.0
            for group in group_source_types
            for level in INCOME_LEVELS
        }
        for group in group_source_types:
            zone_conditional = {
                "below_2.15": conditional[group]["hc215"],
                "2.15_to_6.85": (
                    conditional[group]["hc685"]
                    - conditional[group]["hc215"]
                ),
                "above_6.85": 1.0 - conditional[group]["hc685"],
            }
            for zone_low, zone_high, zone_name in welfare_zones:
                zone_width = zone_high - zone_low
                group_zone_mass = (
                    group_shares[group] * zone_conditional[zone_name]
                )
                if zone_width <= 1e-15:
                    continue
                for income_low, income_high, level in income_intervals:
                    overlap = max(
                        0.0,
                        min(zone_high, income_high)
                        - max(zone_low, income_low),
                    )
                    if overlap > 0:
                        country_joint[(group, level)] += (
                            group_zone_mass * overlap / zone_width
                        )

        for (group, level), value in country_joint.items():
            targets[f"{country} | {group} | {level}"] = (
                country_mass[country] * value
            )
        row_error = max(
            abs(
                sum(country_joint[(group, level)] for level in INCOME_LEVELS)
                - group_shares[group]
            )
            for group in group_source_types
        )
        column_error = max(
            abs(
                sum(
                    country_joint[(group, level)]
                    for group in group_source_types
                )
                - income_shares[level]
            )
            for level in INCOME_LEVELS
        )
        if max(row_error, column_error) > 1e-8:
            raise RuntimeError(
                f"{dimension_name}-income transfer failed margins for {country}: "
                f"row={row_error}, column={column_error}"
            )
        audit[country] = {
            "conditioned_dimension": dimension_name,
            "world_bank_region": region,
            "source_provenance": provenance,
            "source_group_population_shares": source_population_shares,
            "conditioned_margin_shares": group_shares,
            "source_to_conditioned_share_ratios": (
                source_to_margin_share_ratios
            ),
            "direct_group_share_ratio_acceptance": {
                "minimum": MIN_DIRECT_WELFARE_GROUP_SHARE_RATIO,
                "maximum": MAX_DIRECT_WELFARE_GROUP_SHARE_RATIO,
            },
            "country_dimension_direct_accepted": (
                country_dimension_reliable
            ),
            "conditioned_headcounts": conditional,
            "implied_welfare_rank_cuts": {
                "2.15": q215,
                "6.85": q685,
            },
            "maximum_entropy_scope": (
                "within each of the three observed welfare-rank intervals"
            ),
            "maximum_margin_error": max(row_error, column_error),
        }
    return targets, audit


def build_education_income_targets(
    *,
    country_mass: dict[str, float],
    education_targets: dict[str, float],
    income_targets: dict[str, float],
    welfare_rows: Any,
) -> tuple[dict[str, float], dict[str, Any]]:
    conditioned_margins = {
        f"{country} | {group}": value
        for country in COUNTRY_CODES
        for group, value in grouped_education_margins(
            education_targets,
            country,
            WELFARE_EDUCATION_LEVEL_TO_GROUP,
        ).items()
    }
    return build_conditioned_income_targets(
        country_mass=country_mass,
        conditioned_margins=conditioned_margins,
        income_targets=income_targets,
        welfare_rows=welfare_rows,
        group_source_types=WELFARE_EDUCATION_TYPES,
        dimension_name="education",
    )


def build_age_income_targets(
    *,
    country_mass: dict[str, float],
    country_age_sex_targets: dict[str, float],
    income_targets: dict[str, float],
    welfare_rows: Any,
) -> tuple[dict[str, float], dict[str, Any]]:
    conditioned_margins: dict[str, float] = {}
    for country in COUNTRY_CODES:
        youth = sum(
            country_age_sex_targets[
                f"{country} | 18-24 | {sex}"
            ]
            for sex in ("female", "male")
        )
        conditioned_margins[f"{country} | youth"] = youth
        conditioned_margins[f"{country} | adult"] = (
            country_mass[country] - youth
        )
    return build_conditioned_income_targets(
        country_mass=country_mass,
        conditioned_margins=conditioned_margins,
        income_targets=income_targets,
        welfare_rows=welfare_rows,
        group_source_types=WELFARE_AGE_TYPES,
        dimension_name="age_stage",
    )


def direct_ethnicity_distributions(ims: dict[str, Any], wpp: dict[str, Any]) -> dict[str, dict[str, float]]:
    distributions: dict[str, dict[str, float]] = {}

    # UNData table 26, United States 2020.  The official "Asian" count is
    # split using the three Asian origin shares in UN DESA IMS 2020.
    us_asian = 19_886_049
    us_ims = ims["United States"]["origins"]
    east = us_ims.get(906, 0.0)
    south = us_ims.get(5501, 0.0)
    southeast = us_ims.get(920, 0.0)
    asian_total = east + south + southeast
    distributions["United States"] = normalize(
        {
            "White": 204_277_273,
            "Black": 41_104_200,
            "Indigenous": 3_727_135,
            "Pacific Islander": 689_966,
            "Latino": 27_915_715,
            "Mixed": 33_848_943,
            "East Asian": us_asian * east / asian_total,
            "South Asian": us_asian * south / asian_total,
            "SE Asian": us_asian * southeast / asian_total,
            "MENA": 0.0,
        }
    )

    # UNData table 26, Brazil 2022.
    distributions["Brazil"] = normalize(
        {
            "White": 88_252_121,
            "Black": 20_656_458,
            "Mixed": 92_083_286,
            "East Asian": 850_130,
            "Indigenous": 1_227_642,
        }
    )

    # UNData table 26, China 2020.  The national groups do not align with the
    # generator's regional-race taxonomy, so all nationally enumerated groups
    # stay in the broad East Asian bucket; foreigners/unknown remain Mixed.
    distributions["China"] = normalize(
        {"East Asian": 1_408_925_641, "Mixed": 853_083}
    )

    # UNData table 26, United Kingdom 2011.
    distributions["United Kingdom"] = normalize(
        {
            "White": 55_073_552,
            "East Asian": 433_150,
            "South Asian": 1_451_862 + 1_174_983 + 451_529,
            "Black": 1_904_684,
            "Mixed": 861_815 + 1_830_603,
        }
    )

    # Statistics Canada, 2021 Census population group and Indigenous identity.
    canada_total = 36_328_480
    visible_total = 9_639_205
    indigenous = 1_807_250
    distributions["Canada"] = normalize(
        {
            "South Asian": 2_571_400,
            "East Asian": 1_715_770 + 218_140 + 98_890,
            "Black": 1_547_870,
            "SE Asian": 957_355 + 390_340,
            "MENA": 694_015 + 360_495,
            "Latino": 580_235,
            "Mixed": 172_885 + 331_805,
            "Indigenous": indigenous,
            "White": canada_total - visible_total - indigenous,
        }
    )

    # INEGI 2020: 19.4% self-identify as Indigenous and 2.0% as
    # Afro-descendant.  These measures can overlap; normalization makes the
    # mutually-exclusive generator mapping explicit.
    distributions["Mexico"] = normalize(
        {"Indigenous": 0.194, "Black": 0.020, "Latino": 0.806}
    )

    for country, base_ethnicity in BASE_ETHNICITY.items():
        population = wpp["total"][(country, 2020)]
        origins = ims[country]["origins"]
        counts = {ethnicity: 0.0 for ethnicity in ETHNICITIES}
        counts[base_ethnicity] = max(0.0, population - origins[900])
        for origin_code, ethnicity in IMS_ORIGIN_TO_ETHNICITY.items():
            counts[ethnicity] += origins.get(origin_code, 0.0)
        distributions[country] = normalize(counts)
    return distributions


def build_ethnicity_targets(
    country_mass: dict[str, float], distributions: dict[str, dict[str, float]]
) -> dict[str, float]:
    targets: dict[str, float] = {}
    for country, mass in country_mass.items():
        for ethnicity in ETHNICITIES:
            targets[f"{country} | {ethnicity}"] = (
                mass * distributions[country].get(ethnicity, 0.0)
            )
    return targets


def build_occupation_targets(
    wpp: dict[str, Any],
    occupation_rows: Any,
    employment_rows: Any,
) -> tuple[dict[str, float], dict[str, Any]]:
    """Build current-country × sex × occupation from ILOSTAT 2025.

    ILOSTAT supplies employment-to-population ratios and the distribution of
    employed persons by ISCO group.  The complement of the former is the
    explicit ``None`` category in this single-valued profile field.
    """

    sex_codes = {"male": "SEX_M", "female": "SEX_F"}
    adult_population = {
        (country, sex): sum(
            wpp["age_sex"][(country, 2025, age_group, sex)]
            for age_group in AGE_GROUPS
        )
        for country in COUNTRY_CODES
        for sex in sex_codes
    }
    population_total = sum(adult_population.values())
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}

    for country, info in COUNTRY_CODES.items():
        for sex, sex_code in sex_codes.items():
            rate_match = employment_rows[
                (employment_rows["ref_area"] == info["iso3"])
                & (employment_rows["time"] == "2025")
                & (employment_rows["sex"] == sex_code)
                & (employment_rows["classif1"] == "AGE_YTHADULT_YGE15")
            ]
            if len(rate_match) != 1:
                raise RuntimeError(
                    f"Expected one ILOSTAT employment-rate row for {country}/{sex}; "
                    f"found {len(rate_match)}."
                )
            employed_share = min(
                1.0, max(0.0, float(rate_match.iloc[0]["obs_value"]) / 100.0)
            )

            group_counts: dict[str, float] = {}
            for code, label in OCCUPATION_CODES.items():
                match = occupation_rows[
                    (occupation_rows["ref_area"] == info["iso3"])
                    & (occupation_rows["time"] == "2025")
                    & (occupation_rows["sex"] == sex_code)
                    & (occupation_rows["classif1"] == code)
                ]
                if len(match) != 1:
                    raise RuntimeError(
                        f"Expected one ILOSTAT occupation row for "
                        f"{country}/{sex}/{code}; found {len(match)}."
                    )
                group_counts[label] = float(match.iloc[0]["obs_value"])
            employed_distribution = normalize(group_counts)
            population_mass = adult_population[(country, sex)] / population_total
            targets[f"{country} | {sex} | None"] = population_mass * (
                1.0 - employed_share
            )
            for label, share in employed_distribution.items():
                targets[f"{country} | {sex} | {label}"] = (
                    population_mass * employed_share * share
                )
            audit[f"{country} | {sex}"] = {
                "employment_to_population_ratio_age_15_plus": employed_share,
                "occupation_distribution_among_employed": employed_distribution,
                "population_weight_scope": "UN WPP 2025 ages 18-80",
            }
    return targets, audit


def select_ilo_education_rate_panel(
    frame: Any,
    iso3: str,
) -> tuple[dict[tuple[str, str], float], dict[str, Any]] | None:
    subset = frame[
        (frame["ref_area"] == iso3)
        & (frame["sex"].isin(ILO_SEX_CODES))
    ]
    required = {
        (sex, education_group)
        for sex in ILO_SEX_CODES.values()
        for education_group in ILO_EDUCATION_GROUPS
    }
    aggregate_mapping = _education_code_mapping("aggregate")
    subset = subset[subset["classif1"].isin(aggregate_mapping)]
    for year, best, source, candidate in _ilo_candidate_groups(subset):
        values = {
            (
                ILO_SEX_CODES[str(row.sex)],
                aggregate_mapping[str(row.classif1)],
            ): min(1.0, max(0.0, float(row.obs_value) / 100.0))
            for row in candidate.itertuples(index=False)
        }
        if set(values) == required:
            return values, {
                "year": year,
                "source": source,
                "best_source": bool(best),
            }
    return None


def select_ilo_occupation_education_panel(
    frame: Any,
    iso3: str,
) -> tuple[dict[tuple[str, str], float], dict[str, Any]] | None:
    subset = frame[
        (frame["ref_area"] == iso3)
        & (frame["classif1"].isin(ILO_OBSERVED_OCCUPATION_TO_GROUP))
        & (frame["sex"].isin(["SEX_T", *ILO_SEX_CODES]))
    ]
    required = {
        (education_group, occupation_group)
        for education_group in ILO_EDUCATION_GROUPS
        for occupation_group in set(ILO_OBSERVED_OCCUPATION_TO_GROUP.values())
    }
    for year, best, source, candidate in _ilo_candidate_groups(subset):
        for mode in ("aggregate", "isced11", "isced97"):
            education_mapping = _education_code_mapping(mode)
            rows = candidate[candidate["classif2"].isin(education_mapping)]
            for sex_mode in ("total", "sum_sexes"):
                if sex_mode == "total":
                    selected = rows[rows["sex"] == "SEX_T"]
                else:
                    selected = rows[rows["sex"].isin(ILO_SEX_CODES)]
                    if set(selected["sex"]) != set(ILO_SEX_CODES):
                        continue
                values: dict[tuple[str, str], float] = defaultdict(float)
                present: set[tuple[str, str]] = set()
                for row in selected.itertuples(index=False):
                    key = (
                        education_mapping[str(row.classif2)],
                        ILO_OBSERVED_OCCUPATION_TO_GROUP[str(row.classif1)],
                    )
                    values[key] += max(0.0, float(row.obs_value))
                    present.add(key)
                if present == required:
                    return dict(values), {
                        "year": year,
                        "source": source,
                        "best_source": bool(best),
                        "education_classification": mode,
                        "sex_aggregation": sex_mode,
                    }
    return None


def select_ilo_age_stage_occupation_panel(
    frame: Any,
    iso3: str,
) -> tuple[dict[tuple[str, str, str], float], dict[str, Any]] | None:
    subset = frame[
        (frame["ref_area"] == iso3)
        & (frame["sex"].isin(ILO_SEX_CODES))
        & (frame["classif1"].isin(ILO_AGE_STAGE_CODES))
        & (frame["classif2"].isin(ILO_OBSERVED_OCCUPATION_TO_GROUP))
    ]
    required = {
        (stage, sex, occupation)
        for stage in ILO_AGE_STAGE_CODES.values()
        for sex in ILO_SEX_CODES.values()
        for occupation in set(ILO_OBSERVED_OCCUPATION_TO_GROUP.values())
    }
    for year, best, source, candidate in _ilo_candidate_groups(subset):
        values: dict[tuple[str, str, str], float] = defaultdict(float)
        present: set[tuple[str, str, str]] = set()
        for row in candidate.itertuples(index=False):
            value = float(row.obs_value)
            if not math.isfinite(value):
                continue
            key = (
                ILO_AGE_STAGE_CODES[str(row.classif1)],
                ILO_SEX_CODES[str(row.sex)],
                ILO_OBSERVED_OCCUPATION_TO_GROUP[str(row.classif2)],
            )
            values[key] += max(0.0, value)
            present.add(key)
        if present == required:
            return dict(values), {
                "year": year,
                "source": source,
                "best_source": bool(best),
            }
    return None


def select_ilo_age_marital_rate_panel(
    frame: Any,
    iso3: str,
) -> tuple[dict[tuple[str, str, str], float], dict[str, Any]] | None:
    subset = frame[
        (frame["ref_area"] == iso3)
        & (frame["sex"].isin(ILO_SEX_CODES))
        & (frame["classif1"].isin(ILO_AGE_CODES))
        & (
            frame["classif2"].isin(
                ILO_MARITAL_STATUS_TO_RELATIONSHIP
            )
        )
    ]
    for year, best, source, candidate in _ilo_candidate_groups(subset):
        values: dict[tuple[str, str, str], float] = {}
        for row in candidate.itertuples(index=False):
            value = float(row.obs_value)
            if not math.isfinite(value):
                continue
            values[
                (
                    ILO_AGE_CODES[str(row.classif1)],
                    ILO_SEX_CODES[str(row.sex)],
                    ILO_MARITAL_STATUS_TO_RELATIONSHIP[
                        str(row.classif2)
                    ],
                )
            ] = min(1.0, max(0.0, value / 100.0))
        covered_bases = {
            (age_group, sex)
            for age_group, sex, _ in values
        }
        if len(values) >= 24 and len(covered_bases) == 12:
            return values, {
                "year": year,
                "source": source,
                "best_source": bool(best),
                "direct_cells": len(values),
            }
    return None


def build_age_stage_occupation_targets(
    *,
    age_employment_targets: dict[str, float],
    occupation_targets: dict[str, float],
    age_occupation_rows: Any,
) -> tuple[dict[str, float], dict[str, Any]]:
    panels: dict[str, dict[tuple[str, str, str], float]] = {}
    metadata: dict[str, dict[str, Any]] = {}
    global_prior: dict[tuple[str, str, str], float] = defaultdict(float)
    for country, info in COUNTRY_CODES.items():
        selected = select_ilo_age_stage_occupation_panel(
            age_occupation_rows, info["iso3"]
        )
        if selected:
            panel, panel_metadata = selected
            panels[country] = panel
            metadata[country] = panel_metadata
            normalized_panel = normalize(
                {
                    f"{stage} | {sex} | {occupation}": value
                    for (stage, sex, occupation), value in panel.items()
                }
            )
            for key, value in normalized_panel.items():
                stage, sex, occupation = key.split(" | ")
                global_prior[(stage, sex, occupation)] += value
        else:
            metadata[country] = {
                "proxy": "global observed ILO age-stage occupation panel"
            }
    if not global_prior:
        raise RuntimeError("No ILO age-stage occupation panels are available.")

    targets: dict[str, float] = {}
    for country in COUNTRY_CODES:
        panel = panels.get(country, global_prior)
        for sex in ("female", "male"):
            employed_rows = {
                stage: sum(
                    age_employment_targets[
                        f"{country} | {age_group} | {sex} | employed"
                    ]
                    for age_group in AGE_GROUPS
                    if (
                        (stage == "youth" and age_group == "18-24")
                        or (stage == "adult" and age_group != "18-24")
                    )
                )
                for stage in ("youth", "adult")
            }
            not_employed_rows = {
                stage: sum(
                    age_employment_targets[
                        f"{country} | {age_group} | {sex} | not employed"
                    ]
                    for age_group in AGE_GROUPS
                    if (
                        (stage == "youth" and age_group == "18-24")
                        or (stage == "adult" and age_group != "18-24")
                    )
                )
                for stage in ("youth", "adult")
            }
            employed_columns = {
                occupation: occupation_targets[
                    f"{country} | {sex} | {occupation}"
                ]
                for occupation in OCCUPATIONS
                if occupation != "None"
            }
            prior = {
                (stage, occupation): panel.get(
                    (stage, sex, occupation),
                    global_prior.get((stage, sex, occupation), 0.0),
                )
                for stage in ("youth", "adult")
                for occupation in employed_columns
            }
            raked = rake_matrix(
                prior,
                employed_rows,
                employed_columns,
            )
            for (stage, occupation), value in raked.items():
                targets[
                    f"{country} | {stage} | {sex} | {occupation}"
                ] = value
            for stage, value in not_employed_rows.items():
                targets[
                    f"{country} | {stage} | {sex} | None"
                ] = value
    return targets, metadata


def build_relationship_employment_targets(
    *,
    country_relationship_targets: dict[str, float],
    age_employment_targets: dict[str, float],
    age_marital_rate_rows: Any,
) -> tuple[dict[str, float], dict[str, Any]]:
    panels: dict[str, dict[tuple[str, str, str], float]] = {}
    audit: dict[str, Any] = {}
    global_values: dict[tuple[str, str, str], list[float]] = defaultdict(list)
    for country, info in COUNTRY_CODES.items():
        selected = select_ilo_age_marital_rate_panel(
            age_marital_rate_rows, info["iso3"]
        )
        if selected:
            panel, metadata = selected
            panels[country] = panel
            audit[country] = metadata
            for key, value in panel.items():
                global_values[key].append(value)
        else:
            audit[country] = {
                "proxy": "global observed ILO age-sex-marital employment rates"
            }
    global_rates = {
        key: sum(values) / len(values)
        for key, values in global_values.items()
        if values
    }
    if not global_rates:
        raise RuntimeError("No ILO age-marital employment-rate panels are available.")

    targets: dict[str, float] = {}
    for country in COUNTRY_CODES:
        panel = panels.get(country, {})
        fallback_cells = 0
        for age_group in AGE_GROUPS:
            for sex in ("female", "male"):
                relationship_masses = {
                    relationship: country_relationship_targets[
                        f"{country} | {age_group} | {sex} | {relationship}"
                    ]
                    for relationship in RELATIONSHIPS
                }
                base_mass = sum(relationship_masses.values())
                employed_mass = age_employment_targets[
                    f"{country} | {age_group} | {sex} | employed"
                ]
                raw_rates: dict[str, float] = {}
                for relationship in RELATIONSHIPS:
                    key = (age_group, sex, relationship)
                    value = panel.get(key)
                    if value is None and relationship == "in relationship":
                        value = panel.get((age_group, sex, "married"))
                    if value is None:
                        value = global_rates.get(key)
                    if value is None and relationship == "in relationship":
                        value = global_rates.get(
                            (age_group, sex, "married")
                        )
                    if value is None:
                        value = 0.5
                    if key not in panel:
                        fallback_cells += 1
                    raw_rates[relationship] = value
                calibrated = calibrate_probabilities_to_mean(
                    raw_rates,
                    relationship_masses,
                    employed_mass / base_mass,
                )
                for relationship, mass in relationship_masses.items():
                    employed = mass * calibrated[relationship]
                    employed_key = (
                        f"{country} | {relationship} | employed"
                    )
                    not_employed_key = (
                        f"{country} | {relationship} | not employed"
                    )
                    targets[employed_key] = (
                        targets.get(employed_key, 0.0) + employed
                    )
                    targets[not_employed_key] = (
                        targets.get(not_employed_key, 0.0)
                        + mass
                        - employed
                    )
        audit[country]["fallback_cells"] = fallback_cells
    return targets, audit


def build_ilo_joint_targets(
    *,
    wpp: dict[str, Any],
    country_mass: dict[str, float],
    education_targets: dict[str, float],
    occupation_targets: dict[str, float],
    age_rate_rows: Any,
    working_age_education_rows: Any,
    employment_age_education_rows: Any,
    education_rate_rows: Any,
    occupation_education_rows: Any,
) -> tuple[dict[str, dict[str, float]], dict[str, Any]]:
    """Build mutually consistent ILO education/employment/occupation margins."""

    education_age_targets: dict[str, float] = {}
    age_employment_targets: dict[str, float] = {}
    education_employment_targets: dict[str, float] = {}
    education_occupation_targets: dict[str, float] = {}
    audit: dict[str, Any] = {}

    age_education_panels: dict[str, dict[tuple[str, str, str], float]] = {}
    employment_age_education_panels: dict[
        str, dict[tuple[str, str, str], float]
    ] = {}
    occupation_education_panels: dict[
        str, dict[tuple[str, str], float]
    ] = {}
    panel_metadata: dict[str, dict[str, Any]] = {}

    for country, info in COUNTRY_CODES.items():
        age_education = select_ilo_age_education_panel(
            working_age_education_rows, info["iso3"]
        )
        employment_age_education = select_ilo_age_education_panel(
            employment_age_education_rows, info["iso3"]
        )
        occupation_education = select_ilo_occupation_education_panel(
            occupation_education_rows, info["iso3"]
        )
        if age_education:
            age_education_panels[country], age_education_meta = age_education
        else:
            age_education_meta = {"proxy": "independence"}
        if employment_age_education:
            (
                employment_age_education_panels[country],
                employment_age_education_meta,
            ) = employment_age_education
        else:
            employment_age_education_meta = {"proxy": "not available"}
        if occupation_education:
            (
                occupation_education_panels[country],
                occupation_education_meta,
            ) = occupation_education
        else:
            occupation_education_meta = {"proxy": "global observed panel"}
        panel_metadata[country] = {
            "working_age_population_age_education": age_education_meta,
            "employment_age_education": employment_age_education_meta,
            "occupation_education": occupation_education_meta,
        }

    global_occupation_education_prior: dict[tuple[str, str], float] = defaultdict(
        float
    )
    for panel in occupation_education_panels.values():
        normalized_panel = normalize(
            {f"{education} | {occupation}": value for (
                education,
                occupation,
            ), value in panel.items()}
        )
        for key, value in normalized_panel.items():
            education, occupation = key.split(" | ")
            global_occupation_education_prior[(education, occupation)] += value

    education_age_tables: dict[
        str, dict[tuple[str, str, str], float]
    ] = {}
    age_employment_tables: dict[
        str, dict[tuple[str, str, str], float]
    ] = {}
    education_employment_tables: dict[
        str, dict[tuple[str, str, str], float]
    ] = {}

    for country, info in COUNTRY_CODES.items():
        row_targets = {
            f"{age_group} | {sex}": wpp["age_sex"][
                (country, 2025, age_group, sex)
            ]
            / sum(
                wpp["age_sex"][(name, 2025, group, gender)]
                for name in COUNTRY_CODES
                for group in AGE_GROUPS
                for gender in ("female", "male")
            )
            for age_group in AGE_GROUPS
            for sex in ("female", "male")
        }
        education_margins = grouped_education_margins(
            education_targets,
            country,
            ILO_EDUCATION_LEVEL_TO_GROUP,
        )
        panel = age_education_panels.get(country)
        prior = {
            (
                f"{age_group} | {sex}",
                education_group,
            ): (
                panel.get((age_group, sex, education_group), 0.0)
                if panel
                else 1.0
            )
            for age_group in AGE_GROUPS
            for sex in ("female", "male")
            for education_group in ILO_EDUCATION_GROUPS
        }
        raked = rake_matrix(prior, row_targets, education_margins)
        education_age_tables[country] = {
            (
                *row_key.split(" | "),
                education_group,
            ): value
            for (row_key, education_group), value in raked.items()
        }
        for (
            age_group,
            sex,
            education_group,
        ), value in education_age_tables[country].items():
            education_age_targets[
                f"{country} | {age_group} | {sex} | {education_group}"
            ] = value

        age_rate_panel = select_ilo_age_rate_panel(age_rate_rows, info["iso3"])
        if age_rate_panel is None:
            raise RuntimeError(
                f"ILOSTAT has no complete age employment-rate panel for {country}."
            )
        raw_age_rates, age_rate_metadata = age_rate_panel
        for sex in ("female", "male"):
            age_weights = {
                age_group: wpp["age_sex"][
                    (country, 2025, age_group, sex)
                ]
                for age_group in AGE_GROUPS
            }
            country_sex_occupation = {
                occupation: occupation_targets[
                    f"{country} | {sex} | {occupation}"
                ]
                for occupation in OCCUPATIONS
            }
            sex_total = sum(country_sex_occupation.values())
            desired_employed_share = (
                1.0
                - country_sex_occupation["None"] / sex_total
            )
            calibrated_rates = calibrate_probabilities_to_mean(
                {
                    age_group: raw_age_rates[(age_group, sex)]
                    for age_group in AGE_GROUPS
                },
                age_weights,
                desired_employed_share,
            )
            for age_group in AGE_GROUPS:
                mass = row_targets[f"{age_group} | {sex}"]
                employed = mass * calibrated_rates[age_group]
                not_employed = mass - employed
                age_employment_tables.setdefault(country, {})[
                    (age_group, sex, "employed")
                ] = employed
                age_employment_tables[country][
                    (age_group, sex, "not employed")
                ] = not_employed
                age_employment_targets[
                    f"{country} | {age_group} | {sex} | employed"
                ] = employed
                age_employment_targets[
                    f"{country} | {age_group} | {sex} | not employed"
                ] = not_employed

        education_rate_panel = select_ilo_education_rate_panel(
            education_rate_rows, info["iso3"]
        )
        if education_rate_panel:
            raw_education_rates, education_rate_metadata = education_rate_panel
        else:
            working_panel = age_education_panels.get(country)
            employed_panel = employment_age_education_panels.get(country)
            if not working_panel or not employed_panel:
                raw_education_rates = {
                    (sex, education_group): 0.5
                    for sex in ("female", "male")
                    for education_group in ILO_EDUCATION_GROUPS
                }
                education_rate_metadata = {
                    "proxy": "neutral before margin calibration"
                }
            else:
                raw_education_rates = {}
                for sex in ("female", "male"):
                    for education_group in ILO_EDUCATION_GROUPS:
                        working = sum(
                            working_panel[(age_group, sex, education_group)]
                            for age_group in AGE_GROUPS
                        )
                        employed = sum(
                            employed_panel[(age_group, sex, education_group)]
                            for age_group in AGE_GROUPS
                        )
                        raw_education_rates[(sex, education_group)] = (
                            min(1.0, max(0.0, employed / working))
                            if working > 0
                            else 0.5
                        )
                education_rate_metadata = {
                    "derived_from_count_panels": {
                        "working_age": panel_metadata[country][
                            "working_age_population_age_education"
                        ],
                        "employment": panel_metadata[country][
                            "employment_age_education"
                        ],
                    }
                }

        for sex in ("female", "male"):
            sex_education_margins = {
                education_group: sum(
                    education_age_tables[country][
                        (age_group, sex, education_group)
                    ]
                    for age_group in AGE_GROUPS
                )
                for education_group in ILO_EDUCATION_GROUPS
            }
            sex_total = sum(sex_education_margins.values())
            employed_total = sum(
                age_employment_tables[country][
                    (age_group, sex, "employed")
                ]
                for age_group in AGE_GROUPS
            )
            calibrated_education_rates = calibrate_probabilities_to_mean(
                {
                    education_group: raw_education_rates[
                        (sex, education_group)
                    ]
                    for education_group in ILO_EDUCATION_GROUPS
                },
                sex_education_margins,
                employed_total / sex_total,
            )
            for education_group, mass in sex_education_margins.items():
                employed = mass * calibrated_education_rates[education_group]
                not_employed = mass - employed
                education_employment_tables.setdefault(country, {})[
                    (sex, education_group, "employed")
                ] = employed
                education_employment_tables[country][
                    (sex, education_group, "not employed")
                ] = not_employed
                education_employment_targets[
                    f"{country} | {sex} | {education_group} | employed"
                ] = employed
                education_employment_targets[
                    f"{country} | {sex} | {education_group} | not employed"
                ] = not_employed

        education_rows = {
            education_group: sum(
                education_age_tables[country][
                    (age_group, sex, education_group)
                ]
                for age_group in AGE_GROUPS
                for sex in ("female", "male")
            )
            for education_group in ILO_EDUCATION_GROUPS
        }
        occupation_columns = {
            occupation_group: sum(
                occupation_targets[
                    f"{country} | {sex} | {occupation_group}"
                ]
                for sex in ("female", "male")
            )
            for occupation_group in OCCUPATIONS
        }
        observed_occupation_education = occupation_education_panels.get(
            country, global_occupation_education_prior
        )
        not_employed_by_education = {
            education_group: sum(
                education_employment_tables[country][
                    (sex, education_group, "not employed")
                ]
                for sex in ("female", "male")
            )
            for education_group in ILO_EDUCATION_GROUPS
        }
        employed_education_rows = {
            education_group: (
                education_rows[education_group]
                - not_employed_by_education[education_group]
            )
            for education_group in ILO_EDUCATION_GROUPS
        }
        employed_occupation_columns = {
            occupation_group: value
            for occupation_group, value in occupation_columns.items()
            if occupation_group != "None"
        }
        occupation_prior = {}
        for education_group in ILO_EDUCATION_GROUPS:
            for occupation_group in OCCUPATIONS:
                if occupation_group == "None":
                    continue
                occupation_prior[(education_group, occupation_group)] = (
                    observed_occupation_education.get(
                        (education_group, occupation_group), 0.0
                    )
                )
        raked_occupation = rake_matrix(
            occupation_prior,
            employed_education_rows,
            employed_occupation_columns,
        )
        for education_group, value in not_employed_by_education.items():
            education_occupation_targets[
                f"{country} | {education_group} | None"
            ] = value
        for (
            education_group,
            occupation_group,
        ), value in raked_occupation.items():
            education_occupation_targets[
                f"{country} | {education_group} | {occupation_group}"
            ] = value

        audit[country] = {
            **panel_metadata[country],
            "age_employment_rate": age_rate_metadata,
            "education_employment_rate": education_rate_metadata,
            "raking_margins": (
                "UN WPP 2025 age-sex; WDI education; ILOSTAT 2025 modelled "
                "country-sex employment and occupation"
            ),
        }

    return {
        "country_age_sex_education": education_age_targets,
        "country_age_sex_employment": age_employment_targets,
        "country_sex_education_employment": education_employment_targets,
        "country_education_occupation": education_occupation_targets,
    }, audit


def load_ihme_gbd_2023(
    path: Path,
) -> tuple[dict[tuple[str, int, str, int], float], dict[str, Any]]:
    """Load and strictly validate the project's IHME GBD 2023 query export."""

    if not path.exists():
        raise FileNotFoundError(
            f"IHME GBD 2023 CSV not found: {path}. "
            "Pass --ihme-gbd with the verified prevalence export."
        )
    expected_cause_ids = {
        cause_id
        for mapping in (PHYSICAL_CAUSE_IDS, MENTAL_CAUSE_IDS)
        for cause_ids in mapping.values()
        for cause_id in cause_ids
    }
    rates: dict[tuple[str, int, str, int], float] = {}
    cause_names: dict[int, str] = {}
    observed_locations: set[int] = set()
    observed_ages: set[int] = set()
    observed_sexes: set[int] = set()
    observed_causes: set[int] = set()
    row_count = 0
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {
            "population_group_id",
            "measure_id",
            "location_id",
            "sex_id",
            "age_id",
            "cause_id",
            "cause_name",
            "metric_id",
            "year",
            "val",
        }
        missing_columns = required - set(reader.fieldnames or [])
        if missing_columns:
            raise ValueError(
                f"IHME GBD CSV is missing columns: {sorted(missing_columns)}"
            )
        for row in reader:
            row_count += 1
            if (
                int(row["population_group_id"]) != 1
                or int(row["measure_id"]) != 5
                or int(row["metric_id"]) != 3
                or int(row["year"]) != 2023
            ):
                raise ValueError(
                    "IHME GBD CSV must contain all-population prevalence rates "
                    "for year 2023 only."
                )
            location_id = int(row["location_id"])
            age_id = int(row["age_id"])
            sex_id = int(row["sex_id"])
            cause_id = int(row["cause_id"])
            if location_id not in GBD_LOCATION_ID_TO_COUNTRY:
                raise ValueError(f"Unexpected IHME location_id: {location_id}")
            if age_id not in GBD_AGE_ID_TO_RANGE:
                raise ValueError(f"Unexpected IHME age_id: {age_id}")
            if sex_id not in GBD_SEX_ID_TO_SEX:
                raise ValueError(f"Unexpected IHME sex_id: {sex_id}")
            if cause_id not in expected_cause_ids:
                raise ValueError(f"Unexpected IHME cause_id: {cause_id}")
            value = float(row["val"])
            if not math.isfinite(value) or not 0.0 <= value <= 100_000.0:
                raise ValueError(
                    f"Invalid prevalence rate for cause_id={cause_id}: {value}"
                )
            key = (
                GBD_LOCATION_ID_TO_COUNTRY[location_id],
                age_id,
                GBD_SEX_ID_TO_SEX[sex_id],
                cause_id,
            )
            if key in rates:
                raise ValueError(f"Duplicate IHME GBD cell: {key}")
            rates[key] = value
            cause_names[cause_id] = row["cause_name"].strip()
            observed_locations.add(location_id)
            observed_ages.add(age_id)
            observed_sexes.add(sex_id)
            observed_causes.add(cause_id)

    dimensions = (
        ("locations", observed_locations, set(GBD_LOCATION_ID_TO_COUNTRY)),
        ("ages", observed_ages, set(GBD_AGE_ID_TO_RANGE)),
        ("sexes", observed_sexes, set(GBD_SEX_ID_TO_SEX)),
        ("causes", observed_causes, expected_cause_ids),
    )
    for label, observed, expected in dimensions:
        if observed != expected:
            raise ValueError(
                f"IHME GBD {label} mismatch; "
                f"missing={sorted(expected - observed)}, "
                f"unexpected={sorted(observed - expected)}"
            )
    expected_rows = math.prod(
        len(values)
        for values in (
            GBD_LOCATION_ID_TO_COUNTRY,
            GBD_AGE_ID_TO_RANGE,
            GBD_SEX_ID_TO_SEX,
            expected_cause_ids,
        )
    )
    if row_count != expected_rows or len(rates) != expected_rows:
        raise ValueError(
            f"IHME GBD query is incomplete: expected {expected_rows} cells, "
            f"read {row_count} rows and {len(rates)} unique cells."
        )
    return rates, {
        "row_count": row_count,
        "unique_cell_count": len(rates),
        "location_count": len(observed_locations),
        "age_count": len(observed_ages),
        "sex_count": len(observed_sexes),
        "cause_count": len(observed_causes),
        "measure": "Prevalence",
        "metric": "Rate per 100,000",
        "year": 2023,
        "cause_names_by_id": {
            str(cause_id): cause_names[cause_id]
            for cause_id in sorted(cause_names)
        },
    }


def build_one_label_health_targets(
    *,
    rates: dict[tuple[str, int, str, int], float],
    wpp: dict[str, Any],
    country_age_sex_targets: dict[str, float],
    condition_cause_ids: dict[str, tuple[int, ...]],
    condition_statistics: dict[str, Any],
) -> tuple[dict[str, float], dict[str, Any]]:
    """Build country × age × sex × representative-condition targets.

    GBD prevalences are overlapping marginals.  Within every demographic cell,
    the probability of at least one selected condition is approximated with
    ``1-product(1-p_i)`` and allocated across labels in proportion to the raw
    mapped prevalences.  This keeps the source's relative condition rates while
    respecting the schema's one-label representation.
    """

    targets: dict[str, float] = {}
    prevalence_audit: dict[str, dict[str, float]] = {}
    selected_envelopes: list[float] = []
    weighted_raw_prevalence = {
        condition: 0.0 for condition in condition_cause_ids
    }
    for country in COUNTRY_CODES:
        for age_group, (target_low, target_high) in AGE_GROUPS.items():
            for sex in ("female", "male"):
                population_denominator = 0.0
                weighted_rates = {
                    condition: 0.0 for condition in condition_cause_ids
                }
                for age_id, (source_low, source_high) in GBD_AGE_ID_TO_RANGE.items():
                    fraction = overlap_fraction(
                        source_low,
                        source_high,
                        target_low,
                        target_high,
                    )
                    if fraction <= 0:
                        continue
                    population = (
                        wpp["age5_sex"][
                            (country, 2025, source_low, source_high, sex)
                        ]
                        * fraction
                    )
                    population_denominator += population
                    for condition, cause_ids in condition_cause_ids.items():
                        source_rate = sum(
                            rates[(country, age_id, sex, cause_id)]
                            for cause_id in cause_ids
                        )
                        weighted_rates[condition] += source_rate * population
                if population_denominator <= 0:
                    raise RuntimeError(
                        f"No WPP population support for {country}/{age_group}/{sex}"
                    )
                prevalence_per_100k = {
                    condition: value / population_denominator
                    for condition, value in weighted_rates.items()
                }
                probabilities = {
                    condition: min(1.0, max(0.0, value / 100_000.0))
                    for condition, value in prevalence_per_100k.items()
                }
                probability_sum = sum(probabilities.values())
                none_probability = math.prod(
                    1.0 - probability for probability in probabilities.values()
                )
                selected_envelope = 1.0 - none_probability
                selected_envelopes.append(selected_envelope)
                demographic_key = f"{country} | {age_group} | {sex}"
                demographic_mass = country_age_sex_targets[demographic_key]
                targets[f"{demographic_key} | None"] = (
                    demographic_mass * none_probability
                )
                for condition, probability in probabilities.items():
                    allocation = (
                        selected_envelope * probability / probability_sum
                        if probability_sum > 0
                        else 0.0
                    )
                    targets[f"{demographic_key} | {condition}"] = (
                        demographic_mass * allocation
                    )
                    weighted_raw_prevalence[condition] += (
                        demographic_mass * prevalence_per_100k[condition]
                    )
                prevalence_audit[demographic_key] = {
                    condition: round(value, 8)
                    for condition, value in prevalence_per_100k.items()
                }

    return targets, {
        "single_label_semantics": (
            "IHME GBD 2023 cause prevalences are overlapping marginals, while "
            "the profile schema stores one representative condition. Within "
            "each country-age-sex cell, child causes mapped to one schema label "
            "are summed first. The probability of at least one selected label "
            "is then approximated as 1-product(1-p_i), and that envelope is "
            "allocated in proportion to the mapped prevalence rates."
        ),
        "age_aggregation": (
            "IHME five-year age-specific rates are population-weighted with "
            "UN WPP 2025 sex-specific populations. Ages 18-24 use 2/5 of "
            "15-19 plus all of 20-24; ages 65-80 use all of 65-79 plus 1/5 "
            "of 80-84."
        ),
        "condition_statistics": condition_statistics,
        "condition_cause_ids": {
            condition: list(cause_ids)
            for condition, cause_ids in condition_cause_ids.items()
        },
        "population_weighted_raw_prevalence_per_100k": {
            condition: round(value, 8)
            for condition, value in weighted_raw_prevalence.items()
        },
        "population_weighted_selected_condition_share": sum(
            value
            for key, value in targets.items()
            if not key.endswith(" | None")
        ),
        "selected_condition_envelope_range": {
            "minimum": min(selected_envelopes),
            "maximum": max(selected_envelopes),
        },
        "country_age_sex_prevalence_per_100k": prevalence_audit,
    }


def build_health_targets(
    *,
    gbd_path: Path,
    wpp: dict[str, Any],
    country_age_sex_targets: dict[str, float],
) -> tuple[dict[str, float], dict[str, Any], dict[str, float], dict[str, Any]]:
    rates, dataset_audit = load_ihme_gbd_2023(gbd_path)
    physical_targets, physical_audit = build_one_label_health_targets(
        rates=rates,
        wpp=wpp,
        country_age_sex_targets=country_age_sex_targets,
        condition_cause_ids=PHYSICAL_CAUSE_IDS,
        condition_statistics=PHYSICAL_CONDITION_STATISTICS,
    )
    mental_targets, mental_audit = build_one_label_health_targets(
        rates=rates,
        wpp=wpp,
        country_age_sex_targets=country_age_sex_targets,
        condition_cause_ids=MENTAL_CAUSE_IDS,
        condition_statistics=MENTAL_CONDITION_STATISTICS,
    )
    physical_audit["ihme_query_validation"] = dataset_audit
    mental_audit["ihme_query_validation"] = dataset_audit
    return physical_targets, physical_audit, mental_targets, mental_audit


def upsert_constraint(
    data: dict[str, Any],
    *,
    name: str,
    fields: list[str],
    targets: dict[str, float],
) -> None:
    try:
        item = constraint(data, name)
    except StopIteration:
        item = {"name": name, "fields": fields, "targets": {}}
        data["constraints"].append(item)
    item["fields"] = fields
    item["targets"] = rounded(targets)


def rounded(values: dict[str, float]) -> dict[str, float]:
    return {key: round(value, 10) for key, value in values.items()}


def render_values_doc(data: dict[str, Any], path: Path) -> None:
    quality = data["metadata"]["constraint_quality"]
    statuses = data["metadata"]["static_target_status"]
    lines = [
        "# 当前 IPF 约束值",
        "",
        "来源文件：`data/targets/demographic_targets_2025.json`",
        "",
        "本文件由 `scripts/rebuild_official_ipf_targets.py` 自动生成。除非另有说明，"
        f"所有数值均为完整 {len(COUNTRY_CODES)} 国、18–80 岁抽样框中的联合比例。",
        "",
        "## 状态",
        "",
        "| 约束 | 状态 | 来源/处理 |",
        "| --- | --- | --- |",
    ]
    for item in data["constraints"]:
        name = item["name"]
        lines.append(
            f"| `{name}` | `{statuses[name]}` | {quality[name]} |"
        )
    lines.extend(["", "## 数值", ""])
    for item in data["constraints"]:
        name = item["name"]
        lines.extend(
            [
                f"### `{name}`",
                "",
                "字段：" + ", ".join(f"`{field}`" for field in item["fields"]),
                "",
                f"状态：`{statuses[name]}`",
                "",
                f"说明：{quality[name]}",
                "",
                "| 目标 key | 比例 | 百分比 |",
                "| --- | ---: | ---: |",
            ]
        )
        for key, value in item["targets"].items():
            lines.append(f"| `{key}` | {value:.10f} | {value * 100:.6f}% |")
        lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def update_income_constraint_and_metadata(
    data: dict[str, Any],
    income_targets: dict[str, float],
) -> None:
    constraint(data, "country_x_income_level")["targets"] = rounded(income_targets)
    metadata = data["metadata"]
    metadata["title"] = (
        f"PrivacyTrace IPF target constraints for the built-in "
        f"{len(COUNTRY_CODES)}-country generator universe"
    )
    metadata["country_universe"] = list(COUNTRY_CODES)
    metadata["sources"]["income"] = (
        "World Bank Poverty and Inequality Platform (PIP), latest non-interpolated "
        "national observation where available, 2021 PPP welfare distribution. "
        "Bands follow the OECD relative-income classification: below 75%, "
        "75%-100%, 100%-150%, 150%-200%, and at least 200% of the national median."
    )
    metadata["constraint_quality"]["country_x_income_level"] = (
        "source_backed_derived: World Bank PIP 2021 PPP national welfare "
        "distribution; bands relative to the country median"
    )
    metadata["static_target_status"][
        "country_x_income_level"
    ] = "ready_official_source_backed_derived"
    official_rebuild = metadata.setdefault("official_rebuild", {})
    official_rebuild.pop("income_band_definition_2021_ppp_usd_per_day", None)
    official_rebuild["income_band_definition_relative_to_country_median"] = {
        "low": "<0.75x median",
        "lower-middle": "0.75x-1.00x median",
        "middle": "1.00x-1.50x median",
        "upper-middle": "1.50x-2.00x median",
        "high": ">=2.00x median",
    }
    official_rebuild["income_semantics"] = (
        "Relative position in the current country's national household per-capita "
        "welfare distribution. PIP welfare is income for some countries and "
        "consumption for others; it is not individual labor income."
    )


def write_income_audit(
    path: Path,
    income_audit: dict[str, Any],
    education_income_audit: dict[str, Any] | None = None,
    age_income_audit: dict[str, Any] | None = None,
) -> None:
    if path.exists():
        audit = json.loads(path.read_text(encoding="utf-8"))
    else:
        audit = {}
    audit["generated_utc"] = datetime.now(timezone.utc).isoformat()
    audit.setdefault("source_urls", {})[
        "world_bank_pip_api"
    ] = "https://api.worldbank.org/pip/v1"
    audit["source_urls"]["oecd_relative_income_classification"] = (
        OECD_MIDDLE_INCOME_URL
    )
    if education_income_audit is not None:
        audit["source_urls"]["world_bank_global_welfare_profile_2022"] = (
            WORLD_BANK_WELFARE_PROFILE_URL
        )
        audit["education_income_joint"] = education_income_audit
    if age_income_audit is not None:
        audit["source_urls"]["world_bank_global_welfare_profile_2022"] = (
            WORLD_BANK_WELFARE_PROFILE_URL
        )
        audit["age_income_joint"] = age_income_audit
    audit["income"] = income_audit
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--target", type=Path, default=DEFAULT_TARGET)
    parser.add_argument("--audit", type=Path, default=DEFAULT_AUDIT)
    parser.add_argument("--values-doc", type=Path, default=DEFAULT_VALUES_DOC)
    parser.add_argument("--cache-dir", type=Path, default=DEFAULT_CACHE)
    parser.add_argument(
        "--city-catalog",
        type=Path,
        default=DEFAULT_CITY_CATALOG,
    )
    parser.add_argument(
        "--ihme-gbd",
        type=Path,
        default=DEFAULT_IHME_GBD_2023,
        help=(
            "Verified IHME GBD 2023 CSV containing prevalence rates for the "
            "configured countries, sexes, ages, and causes."
        ),
    )
    parser.add_argument(
        "--income-only",
        action="store_true",
        help=(
            "Rebuild country_x_income_level plus its dependent "
            "education-by-income and age-stage-by-income joint constraints "
            "and audit metadata."
        ),
    )
    args = parser.parse_args()

    args.cache_dir.mkdir(parents=True, exist_ok=True)
    data = json.loads(args.target.read_text(encoding="utf-8"))
    obsolete_constraints = {
        "current_country_x_current_city",
        "physical_condition",
        "sex_x_mental_condition",
        (
            "current_country_x_age_group_x_sex_x_relationship_status"
            "_x_employment_status"
        ),
    }
    data["constraints"] = [
        item
        for item in data["constraints"]
        if item["name"] not in obsolete_constraints
    ]
    for metadata_key in ("constraint_quality", "static_target_status"):
        for name in obsolete_constraints:
            data["metadata"].get(metadata_key, {}).pop(name, None)
    if args.income_only:
        country_mass = constraint(data, "country_population")["targets"]
        pip_rows = fetch_pip_rows(args.cache_dir)
        income_targets, income_audit = build_income_targets(country_mass, pip_rows)
        welfare_rows = load_world_bank_welfare_profile(
            download(
                WORLD_BANK_WELFARE_PROFILE_URL,
                args.cache_dir / "world_bank_pppr_2024_reproducibility.zip",
            )
        )
        education_income_targets, education_income_audit = (
            build_education_income_targets(
                country_mass=country_mass,
                education_targets=constraint(
                    data, "country_x_education"
                )["targets"],
                income_targets=income_targets,
                welfare_rows=welfare_rows,
            )
        )
        age_income_targets, age_income_audit = build_age_income_targets(
            country_mass=country_mass,
            country_age_sex_targets=constraint(
                data, "current_country_x_age_group_x_sex"
            )["targets"],
            income_targets=income_targets,
            welfare_rows=welfare_rows,
        )
        update_income_constraint_and_metadata(data, income_targets)
        upsert_constraint(
            data,
            name=(
                "current_country_x_welfare_education_group_x_income_level"
            ),
            fields=[
                "current_country",
                "welfare_education_group",
                "income_level",
            ],
            targets=education_income_targets,
        )
        upsert_constraint(
            data,
            name="current_country_x_age_stage_x_income_level",
            fields=[
                "current_country",
                "age_stage",
                "income_level",
            ],
            targets=age_income_targets,
        )
        metadata = data["metadata"]
        metadata["sources"]["education_income_joint"] = (
            "World Bank Global distribution of welfare 2022 education-"
            "conditioned $2.15/$6.85 poverty rates, combined with WDI "
            "education and PIP median-relative income margins."
        )
        metadata["constraint_quality"][
            "current_country_x_welfare_education_group_x_income_level"
        ] = (
            "world_bank_source_backed_derived: education-conditioned "
            "$2.15/$6.85 welfare-rank dependence transferred by maximum "
            "entropy to PIP median-relative income bands"
        )
        metadata["constraint_quality"][
            "current_country_x_age_stage_x_income_level"
        ] = (
            "world_bank_source_backed_derived: youth/adult-conditioned "
            "$2.15/$6.85 welfare-rank dependence transferred by maximum "
            "entropy to PIP median-relative income bands"
        )
        metadata["static_target_status"][
            "current_country_x_welfare_education_group_x_income_level"
        ] = "ready_world_bank_source_backed_maximum_entropy"
        metadata["static_target_status"][
            "current_country_x_age_stage_x_income_level"
        ] = "ready_world_bank_source_backed_maximum_entropy"
        metadata.setdefault("schema_option_set", {})[
            "welfare_education_group"
        ] = list(WELFARE_EDUCATION_TYPES)
        metadata["schema_option_set"]["age_stage"] = list(
            WELFARE_AGE_TYPES
        )
        metadata.setdefault("official_rebuild", {})[
            "education_income_joint_inputs"
        ] = education_income_audit
        metadata["official_rebuild"][
            "age_income_joint_inputs"
        ] = age_income_audit
        data["metadata"]["created_utc"] = datetime.now(timezone.utc).date().isoformat()
        data["metadata"]["official_rebuild"][
            "generated_utc"
        ] = datetime.now(timezone.utc).isoformat()
        args.target.write_text(
            json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
        render_values_doc(data, args.values_doc)
        write_income_audit(
            args.audit,
            income_audit,
            education_income_audit,
            age_income_audit,
        )
        return

    wpp = load_wpp(download(WPP_URL, args.cache_dir / "wpp2024_age5_sex.csv.gz"))
    country_mass, age_sex_mass = build_population_targets(wpp)
    city_catalog, city_size_targets = load_wup_city_catalog(
        download(
            WUP_CITY_URL,
            args.cache_dir / "wup2025_cities_population.xlsx",
        ),
        country_mass,
    )
    education_rows = fetch_wdi_education(args.cache_dir)
    education_targets, education_audit = build_education_targets(
        country_mass,
        education_rows,
    )
    religion_targets, religion_audit = load_pew_religion(
        download(
            PEW_RELIGION_URL,
            args.cache_dir / "pew_religious_composition_2010_2020.zip",
        ),
        country_mass,
    )
    ims = load_ims(download(IMS_URL, args.cache_dir / "ims2020_destination_origin.xlsx"))
    ilo_headers = {
        "User-Agent": "R/4.5 Rilostat/2.5.0",
        "Accept-Encoding": "gzip",
        "Referer": "https://rplumber.ilo.org/",
    }
    occupation_rows = load_ilo_rds(
        download(
            ILO_OCCUPATION_URL,
            args.cache_dir / "ilo_emp_occupation_2025.rds",
            headers=ilo_headers,
        )
    )
    employment_rows = load_ilo_rds(
        download(
            ILO_EMPLOYMENT_RATE_URL,
            args.cache_dir / "ilo_employment_rate_2025.rds",
            headers=ilo_headers,
        )
    )
    age_employment_rate_rows = load_ilo_rds(
        download(
            ILO_EMPLOYMENT_RATE_BY_AGE_URL,
            args.cache_dir / "ilo_employment_rate_by_age_observed.rds",
            headers=ilo_headers,
        )
    )
    working_age_education_rows = load_ilo_rds(
        download(
            ILO_WORKING_AGE_AGE_EDUCATION_URL,
            args.cache_dir / "ilo_working_age_population_by_age_education.rds",
            headers=ilo_headers,
        )
    )
    employment_age_education_rows = load_ilo_rds(
        download(
            ILO_EMPLOYMENT_AGE_EDUCATION_URL,
            args.cache_dir / "ilo_employment_by_age_education.rds",
            headers=ilo_headers,
        )
    )
    education_employment_rate_rows = load_ilo_rds(
        download(
            ILO_EMPLOYMENT_RATE_BY_EDUCATION_URL,
            args.cache_dir / "ilo_employment_rate_by_education.rds",
            headers=ilo_headers,
        )
    )
    occupation_education_rows = load_ilo_rds(
        download(
            ILO_OCCUPATION_EDUCATION_URL,
            args.cache_dir / "ilo_employment_by_occupation_education.rds",
            headers=ilo_headers,
        )
    )
    age_occupation_rows = load_ilo_rds(
        download(
            ILO_AGE_OCCUPATION_URL,
            args.cache_dir / "ilo_employment_by_age_occupation.rds",
            headers=ilo_headers,
        )
    )
    age_marital_employment_rate_rows = load_ilo_rds(
        download(
            ILO_AGE_MARITAL_EMPLOYMENT_RATE_URL,
            args.cache_dir
            / "ilo_employment_rate_by_age_marital_status.rds",
            headers=ilo_headers,
        )
    )
    migration_targets, migration_audit = build_migration_targets(country_mass, wpp, ims)
    migration_status_targets, migration_status_audit = (
        build_migration_status_by_sex_targets(wpp, ims)
    )
    marital_rates, marital_years = load_marital_rows(args.cache_dir)
    relationship_targets = build_relationship_targets(age_sex_mass, wpp, marital_rates)
    country_relationship_targets = build_country_relationship_targets(
        wpp, marital_rates
    )
    pip_rows = fetch_pip_rows(args.cache_dir)
    income_targets, income_audit = build_income_targets(country_mass, pip_rows)
    ethnicity_distributions = direct_ethnicity_distributions(ims, wpp)
    ethnicity_targets = build_ethnicity_targets(country_mass, ethnicity_distributions)
    occupation_targets, occupation_audit = build_occupation_targets(
        wpp, occupation_rows, employment_rows
    )
    country_age_sex_targets = build_country_age_sex_targets(wpp)
    ilo_joint_targets, ilo_joint_audit = build_ilo_joint_targets(
        wpp=wpp,
        country_mass=country_mass,
        education_targets=education_targets,
        occupation_targets=occupation_targets,
        age_rate_rows=age_employment_rate_rows,
        working_age_education_rows=working_age_education_rows,
        employment_age_education_rows=employment_age_education_rows,
        education_rate_rows=education_employment_rate_rows,
        occupation_education_rows=occupation_education_rows,
    )
    age_stage_occupation_targets, age_stage_occupation_audit = (
        build_age_stage_occupation_targets(
            age_employment_targets=ilo_joint_targets[
                "country_age_sex_employment"
            ],
            occupation_targets=occupation_targets,
            age_occupation_rows=age_occupation_rows,
        )
    )
    relationship_employment_targets, relationship_employment_audit = (
        build_relationship_employment_targets(
            country_relationship_targets=country_relationship_targets,
            age_employment_targets=ilo_joint_targets[
                "country_age_sex_employment"
            ],
            age_marital_rate_rows=age_marital_employment_rate_rows,
        )
    )
    welfare_rows = load_world_bank_welfare_profile(
        download(
            WORLD_BANK_WELFARE_PROFILE_URL,
            args.cache_dir / "world_bank_pppr_2024_reproducibility.zip",
        )
    )
    education_income_targets, education_income_audit = (
        build_education_income_targets(
            country_mass=country_mass,
            education_targets=education_targets,
            income_targets=income_targets,
            welfare_rows=welfare_rows,
        )
    )
    age_income_targets, age_income_audit = build_age_income_targets(
        country_mass=country_mass,
        country_age_sex_targets=country_age_sex_targets,
        income_targets=income_targets,
        welfare_rows=welfare_rows,
    )
    (
        physical_targets,
        physical_audit,
        mental_targets,
        mental_audit,
    ) = build_health_targets(
        gbd_path=args.ihme_gbd,
        wpp=wpp,
        country_age_sex_targets=country_age_sex_targets,
    )

    constraint(data, "country_population")["targets"] = rounded(country_mass)
    constraint(data, "age_x_sex")["targets"] = rounded(age_sex_mass)
    upsert_constraint(
        data,
        name="current_country_x_age_group_x_sex",
        fields=["current_country", "age_group", "sex"],
        targets=country_age_sex_targets,
    )
    constraint(data, "country_x_education")["targets"] = rounded(
        education_targets
    )
    constraint(data, "country_x_religion")["targets"] = rounded(
        religion_targets
    )
    upsert_constraint(
        data,
        name="current_country_x_current_city_size",
        fields=["current_country", "current_city_size"],
        targets=city_size_targets,
    )
    constraint(data, "current_country_x_citizenship_country")["targets"] = rounded(
        migration_targets
    )
    constraint(data, "current_country_x_birth_country")["targets"] = rounded(
        migration_targets
    )
    upsert_constraint(
        data,
        name="current_country_x_sex_x_birth_migration_status",
        fields=[
            "current_country",
            "sex",
            "birth_migration_status",
        ],
        targets=migration_status_targets,
    )
    upsert_constraint(
        data,
        name="current_country_x_sex_x_citizenship_migration_status",
        fields=[
            "current_country",
            "sex",
            "citizenship_migration_status",
        ],
        targets=migration_status_targets,
    )
    try:
        old_ethnicity = constraint(data, "citizenship_country_x_ethnicity")
    except StopIteration:
        old_ethnicity = constraint(data, "current_country_x_ethnicity")
    old_ethnicity["name"] = "current_country_x_ethnicity"
    old_ethnicity["fields"] = ["current_country", "ethnicity"]
    old_ethnicity["targets"] = rounded(ethnicity_targets)
    update_income_constraint_and_metadata(data, income_targets)
    constraint(data, "age_x_sex_x_relationship_status")["targets"] = rounded(
        relationship_targets
    )
    upsert_constraint(
        data,
        name=(
            "current_country_x_age_group_x_sex_x_relationship_status"
        ),
        fields=[
            "current_country",
            "age_group",
            "sex",
            "relationship_status",
        ],
        targets=country_relationship_targets,
    )
    upsert_constraint(
        data,
        name=(
            "current_country_x_relationship_status_x_employment_status"
        ),
        fields=[
            "current_country",
            "relationship_status",
            "employment_status",
        ],
        targets=relationship_employment_targets,
    )
    upsert_constraint(
        data,
        name="current_country_x_sex_x_occupation",
        fields=["current_country", "sex", "occupation_group"],
        targets=occupation_targets,
    )
    upsert_constraint(
        data,
        name="current_country_x_age_group_x_sex_x_ilo_education_group",
        fields=[
            "current_country",
            "age_group",
            "sex",
            "ilo_education_group",
        ],
        targets=ilo_joint_targets["country_age_sex_education"],
    )
    upsert_constraint(
        data,
        name="current_country_x_age_group_x_sex_x_employment_status",
        fields=[
            "current_country",
            "age_group",
            "sex",
            "employment_status",
        ],
        targets=ilo_joint_targets["country_age_sex_employment"],
    )
    upsert_constraint(
        data,
        name="current_country_x_sex_x_ilo_education_group_x_employment_status",
        fields=[
            "current_country",
            "sex",
            "ilo_education_group",
            "employment_status",
        ],
        targets=ilo_joint_targets[
            "country_sex_education_employment"
        ],
    )
    upsert_constraint(
        data,
        name="current_country_x_ilo_education_group_x_occupation_group",
        fields=[
            "current_country",
            "ilo_education_group",
            "occupation_group",
        ],
        targets=ilo_joint_targets["country_education_occupation"],
    )
    upsert_constraint(
        data,
        name=(
            "current_country_x_age_stage_x_sex_x_occupation_group"
        ),
        fields=[
            "current_country",
            "age_stage",
            "sex",
            "occupation_group",
        ],
        targets=age_stage_occupation_targets,
    )
    upsert_constraint(
        data,
        name="current_country_x_welfare_education_group_x_income_level",
        fields=[
            "current_country",
            "welfare_education_group",
            "income_level",
        ],
        targets=education_income_targets,
    )
    upsert_constraint(
        data,
        name="current_country_x_age_stage_x_income_level",
        fields=[
            "current_country",
            "age_stage",
            "income_level",
        ],
        targets=age_income_targets,
    )
    upsert_constraint(
        data,
        name=(
            "current_country_x_age_group_x_sex_x_physical_condition"
        ),
        fields=[
            "current_country",
            "age_group",
            "sex",
            "physical_condition",
        ],
        targets=physical_targets,
    )
    upsert_constraint(
        data,
        name=(
            "current_country_x_age_group_x_sex_x_mental_condition"
        ),
        fields=[
            "current_country",
            "age_group",
            "sex",
            "mental_condition",
        ],
        targets=mental_targets,
    )

    metadata = data["metadata"]
    metadata["created_utc"] = datetime.now(timezone.utc).date().isoformat()
    metadata["important_caveat"] = (
        "All IPF constraints now use official or intergovernmental statistical inputs, "
        "except country_x_religion, which remains authority-backed Pew data. Health "
        "targets use an IHME GBD 2023 prevalence query by country, age and sex. "
        "Derived constraints retain their transformations and proxy semantics in "
        "official_rebuild."
    )
    metadata["sources"].update(
        {
            "current_city": (
                "UN DESA World Urbanization Prospects 2025 File 21, DEGURBA "
                "cities with at least 50,000 inhabitants in 2025."
            ),
            "education": (
                "World Bank API, WDI educational-attainment cumulative "
                "indicators for population age 25+, latest non-null observation."
            ),
            "education_employment_occupation_joint": (
                "ILOSTAT annual observed working-age population by sex, age "
                "and education; employment-to-population ratios by sex, age "
                "or education; and employment by sex, occupation and education."
            ),
            "age_occupation_joint": (
                "ILOSTAT observed employment by sex, youth/adult age stage "
                "and ISCO-08 occupation, raked to the modelled 2025 "
                "employment and occupation margins."
            ),
            "relationship_employment_joint": (
                "ILOSTAT observed employment-to-population ratios by sex, "
                "10-year age band and marital status, calibrated to the "
                "UNData relationship and ILO 2025 employment margins."
            ),
            "education_income_joint": (
                "World Bank Global distribution of welfare 2022 education-"
                "conditioned $2.15/$6.85 poverty rates, combined with WDI "
                "education and PIP median-relative income margins."
            ),
            "age_income_joint": (
                "World Bank Global distribution of welfare 2022 youth/adult-"
                "conditioned $2.15/$6.85 poverty rates, combined with UN WPP "
                "age-stage and PIP median-relative income margins."
            ),
            "religion": (
                "Pew Research Center Religious Composition by Country, "
                "2010-2020 dataset, 2020 percentages."
            ),
            "relationship_status": (
                "UNSD Demographic Statistics Database / UNData table 23, latest "
                "available national observation by country, age and sex."
            ),
            "ethnicity": (
                "UNData table 26 and official national census population-group "
                "tables; UN DESA International Migrant Stock 2020 regional-origin "
                "mapping where no compatible national ethnicity table exists."
            ),
            "birth_country": (
                "UN DESA International Migrant Stock 2020 destination-origin "
                f"matrix, conditioned on the generator's supported "
                f"{len(COUNTRY_CODES)}-country universe."
            ),
            "citizenship_country": (
                "UN DESA International Migrant Stock 2020 destination-origin "
                f"matrix, conditioned on the generator's supported "
                f"{len(COUNTRY_CODES)}-country universe."
            ),
            "occupation": (
                "ILOSTAT modelled estimates, 2025: employment by sex and ISCO-08 "
                "occupation (EMP_2EMP_SEX_OCU_NB_A) plus employment-to-population "
                "ratio by sex (EMP_2WAP_SEX_AGE_RT_A)."
            ),
            "physical_condition": (
                "IHME Global Burden of Disease 2023 prevalence rates for 26 "
                "countries, two sexes and five-year age groups, mapped to 22 "
                "physical-condition labels and population-weighted with UN WPP 2025."
            ),
            "mental_condition": (
                "IHME Global Burden of Disease 2023 prevalence rates for 26 "
                "countries, two sexes and five-year age groups, mapped to 16 "
                "mental/substance-use labels and population-weighted with UN WPP 2025."
            ),
        }
    )
    metadata["constraint_quality"].update(
        {
            "current_country_x_citizenship_country": (
                "official_derived: UN DESA IMS 2020; direct where destination data "
                "type contains C, official migration-origin proxy where it contains B"
            ),
            "current_country_x_birth_country": (
                "official_derived: UN DESA IMS 2020; direct where destination data "
                "type contains B, official migration-origin proxy where it contains C"
            ),
            "current_country_x_sex_x_birth_migration_status": (
                "official_derived: UN DESA IMS 2020 destination-origin-sex "
                "matrix, reduced to local/foreign status within supported countries"
            ),
            "current_country_x_sex_x_citizenship_migration_status": (
                "official_derived_proxy: same IMS 2020 sex-specific matrix; "
                "direct where destination basis includes citizenship, birth proxy otherwise"
            ),
            "current_country_x_ethnicity": (
                "official_derived: UNData table 26, national censuses, and documented "
                "regional-origin mapping for countries without compatible ethnicity tables"
            ),
            "age_x_sex_x_relationship_status": (
                "source_backed: latest UNData table 23 marital-status observations "
                "combined with UN WPP 2025 age-sex weights"
            ),
            "current_country_x_age_group_x_sex": (
                "official_modelled: UN WPP 2024 medium projection for 2025, "
                "ages 18-80"
            ),
            "current_country_x_age_group_x_sex_x_relationship_status": (
                "official_source_backed: country-specific UNData table 23 "
                "marital-status rates combined with UN WPP 2025 age-sex weights"
            ),
            "current_country_x_relationship_status_x_employment_status": (
                "official_raked: ILOSTAT observed age-sex-marital employment "
                "rates logit-calibrated to consistent UN relationship and "
                "ILO employment margins, then aggregated by country and "
                "relationship status"
            ),
            "current_country_x_sex_x_occupation": (
                "official_modelled_derived: ILOSTAT 2025 employment ratio and "
                "employment by ISCO-08 group, combined with UN WPP adult weights"
            ),
            "current_country_x_age_group_x_sex_x_ilo_education_group": (
                "official_raked: ILOSTAT observed age-sex-education structure "
                "raked to WPP 2025 age-sex and WDI education margins"
            ),
            "current_country_x_age_group_x_sex_x_employment_status": (
                "official_raked: ILOSTAT observed 10-year age employment curves "
                "logit-calibrated to 2025 modelled country-sex employment margins"
            ),
            "current_country_x_sex_x_ilo_education_group_x_employment_status": (
                "official_raked: ILOSTAT education-specific employment rates "
                "calibrated to the consistent education and employment margins"
            ),
            "current_country_x_ilo_education_group_x_occupation_group": (
                "official_raked: ILOSTAT observed occupation by education, "
                "raked to consistent country education and occupation margins"
            ),
            "current_country_x_age_stage_x_sex_x_occupation_group": (
                "official_raked: ILOSTAT observed youth/adult occupation by "
                "sex, raked to consistent age-employment and occupation margins"
            ),
            "current_country_x_welfare_education_group_x_income_level": (
                "world_bank_source_backed_derived: education-conditioned "
                "$2.15/$6.85 welfare-rank dependence transferred by maximum "
                "entropy to PIP median-relative income bands"
            ),
            "current_country_x_age_stage_x_income_level": (
                "world_bank_source_backed_derived: youth/adult-conditioned "
                "$2.15/$6.85 welfare-rank dependence transferred by maximum "
                "entropy to PIP median-relative income bands"
            ),
            "current_country_x_current_city_size": (
                "official_derived: UN WUP 2025 city populations aggregated "
                "to seven population-size classes within each country"
            ),
            "current_country_x_age_group_x_sex_x_physical_condition": (
                "ihme_gbd_2023_country_age_sex_derived: cause-specific prevalence "
                "rates population-weighted to project age bands and converted to "
                "one representative physical label with an explicit overlap approximation"
            ),
            "current_country_x_age_group_x_sex_x_mental_condition": (
                "ihme_gbd_2023_country_age_sex_derived: cause-specific prevalence "
                "rates population-weighted to project age bands and converted to "
                "one representative mental/substance-use label"
            ),
        }
    )
    metadata["constraint_quality"].pop("citizenship_country_x_ethnicity", None)
    metadata["static_target_status"].update(
        {
            "current_country_x_citizenship_country": "ready_official_derived_proxy_mixed_basis",
            "current_country_x_birth_country": "ready_official_derived_proxy_mixed_basis",
            "current_country_x_sex_x_birth_migration_status": "ready_official_derived",
            "current_country_x_sex_x_citizenship_migration_status": "ready_official_derived_proxy_mixed_basis",
            "current_country_x_ethnicity": "ready_official_derived",
            "age_x_sex_x_relationship_status": "ready_official_source_backed",
            "current_country_x_age_group_x_sex": "ready_official_modelled",
            "current_country_x_age_group_x_sex_x_relationship_status": "ready_official_source_backed",
            "current_country_x_relationship_status_x_employment_status": "ready_official_raked",
            "current_country_x_sex_x_occupation": "ready_official_modelled_derived",
            "current_country_x_age_group_x_sex_x_ilo_education_group": "ready_official_raked",
            "current_country_x_age_group_x_sex_x_employment_status": "ready_official_raked",
            "current_country_x_sex_x_ilo_education_group_x_employment_status": "ready_official_raked",
            "current_country_x_ilo_education_group_x_occupation_group": "ready_official_raked",
            "current_country_x_age_stage_x_sex_x_occupation_group": "ready_official_raked",
            "current_country_x_welfare_education_group_x_income_level": "ready_world_bank_source_backed_maximum_entropy",
            "current_country_x_age_stage_x_income_level": "ready_world_bank_source_backed_maximum_entropy",
            "current_country_x_current_city_size": "ready_official_derived",
            "current_country_x_age_group_x_sex_x_physical_condition": (
                "ready_ihme_gbd_2023_country_age_sex_derived"
            ),
            "current_country_x_age_group_x_sex_x_mental_condition": (
                "ready_ihme_gbd_2023_country_age_sex_derived"
            ),
        }
    )
    metadata["static_target_status"].pop("citizenship_country_x_ethnicity", None)
    metadata["replacement_notes"]["added"] = [
        "current_country_x_citizenship_country",
        "current_country_x_birth_country",
        "current_country_x_sex_x_birth_migration_status",
        "current_country_x_sex_x_citizenship_migration_status",
        "current_country_x_ethnicity",
        "country_x_income_level",
        "current_country_x_age_group_x_sex",
        "age_x_sex_x_relationship_status",
        "current_country_x_age_group_x_sex_x_relationship_status",
        "current_country_x_relationship_status_x_employment_status",
        "current_country_x_current_city_size",
        "current_country_x_sex_x_occupation",
        "current_country_x_age_group_x_sex_x_ilo_education_group",
        "current_country_x_age_group_x_sex_x_employment_status",
        "current_country_x_sex_x_ilo_education_group_x_employment_status",
        "current_country_x_ilo_education_group_x_occupation_group",
        "current_country_x_age_stage_x_sex_x_occupation_group",
        "current_country_x_welfare_education_group_x_income_level",
        "current_country_x_age_stage_x_income_level",
        "current_country_x_age_group_x_sex_x_physical_condition",
        "current_country_x_age_group_x_sex_x_mental_condition",
    ]
    project_root_text = str(PROJECT_ROOT)
    if project_root_text not in sys.path:
        sys.path.insert(0, project_root_text)
    from privacy_trace.profile_generator import (
        MENTAL_CONDITIONS as SCHEMA_MENTAL_CONDITIONS,
        OCCUPATIONS as DETAILED_OCCUPATIONS,
        PHYSICAL_CONDITIONS as SCHEMA_PHYSICAL_CONDITIONS,
    )

    metadata.setdefault("schema_option_set", {}).update(
        {
            "citizenship": list(COUNTRY_CODES),
            "current_country": list(COUNTRY_CODES),
            "birth_country": list(COUNTRY_CODES),
            "occupation": DETAILED_OCCUPATIONS,
            "occupation_group": OCCUPATIONS,
            "employment_status": ["employed", "not employed"],
            "age_stage": list(WELFARE_AGE_TYPES),
            "ilo_education_group": list(ILO_EDUCATION_GROUPS),
            "welfare_education_group": list(WELFARE_EDUCATION_TYPES),
            "birth_migration_status": ["local", "foreign"],
            "citizenship_migration_status": ["local", "foreign"],
            "physical_condition": SCHEMA_PHYSICAL_CONDITIONS,
            "mental_condition": SCHEMA_MENTAL_CONDITIONS,
        }
    )
    metadata["schema_option_set"].pop("physical_condition_group", None)
    metadata["schema_option_set"].pop("mental_condition_group", None)
    metadata["official_rebuild"] = {
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "migration_universe": (
            "For birth and citizenship country, each destination row is conditioned "
            f"on native/same-country plus foreign origins among the {len(COUNTRY_CODES)} countries "
            "representable by the current generator."
        ),
        "age_scope_caveat": (
            "Migration, income, and some ethnicity inputs have no consistent age "
            "18-80 slice. Their conditional distributions are applied to the adult "
            "frame while WPP controls the adult country and age-sex marginals."
        ),
        "migration_basis_by_destination": migration_audit,
        "migration_status_by_sex": migration_status_audit,
        "income_band_definition_relative_to_country_median": metadata[
            "official_rebuild"
        ]["income_band_definition_relative_to_country_median"],
        "income_semantics": metadata["official_rebuild"]["income_semantics"],
        "relationship_mapping": {
            "single": ["Single (never married)"],
            "in relationship": ["In consensual union"],
            "married": ["Married"],
            "divorced": ["Divorced and not remarried", "Married but separated"],
            "widowed": ["Widowed and not remarried"],
        },
        "relationship_source_years": marital_years,
        "ilo_joint_constraint_inputs": ilo_joint_audit,
        "ilo_age_stage_occupation_inputs": age_stage_occupation_audit,
        "ilo_relationship_employment_inputs": relationship_employment_audit,
        "education_income_joint_inputs": education_income_audit,
        "age_income_joint_inputs": age_income_audit,
        "education_income_semantics": (
            "Four World Bank education groups are crossed with the project's "
            "five PIP median-relative welfare bands. Observed education-specific "
            "headcounts at $2.15 and $6.85 define three welfare-rank intervals; "
            "maximum entropy is used only inside those intervals. Country WDI "
            "education and PIP income margins remain exact. Missing country "
            "education rows use the World Bank OHI regional relation and are "
            "flagged per country."
        ),
        "age_income_semantics": (
            "World Bank youth [15-24] and adult [25+] welfare profiles are "
            "mapped to the generator's age_stage youth [18-24] and adult "
            "[25-80]. The source's youth band therefore includes ages 15-17 "
            "that are outside the generator. Country age-stage and PIP income "
            "margins remain exact; missing country rows use the OHI relation "
            "and are flagged per country."
        ),
        "ethnicity_conditioning_change": (
            "Changed from citizenship_country x ethnicity to current_country x "
            "ethnicity because official census population-group tables describe "
            "resident populations, not citizens abroad."
        ),
        "occupation_semantics": (
            "None means not employed under the ILOSTAT employment-to-population "
            "measure. Employed profiles use the eight ISCO-08 groups available in "
            "the modelled 2025 bulk table, then receive one of 115 concise concrete "
            "job labels mapped to that group. These display labels are not "
            "independently calibrated. The ILO age-15-plus employment ratio is "
            "applied to the generator's WPP-weighted age-18-to-80 frame."
        ),
        "age_stage_occupation_semantics": (
            "ILOSTAT youth [15-24] and adult [25+] occupation counts are "
            "mapped to generator ages 18-24 and 25-80. Observed dependence "
            "is raked to the existing country-sex occupation and age-stage "
            "employment margins; China and Japan use the pooled observed "
            "ILOSTAT dependence and are explicitly flagged."
        ),
        "relationship_employment_semantics": (
            "ILOSTAT marital-status employment-rate shapes are calibrated "
            "within every country-age-sex cell to the UNData relationship "
            "shares and the ILO employment margin. Missing consensual-union "
            "cells use the same-country married rate first, then the pooled "
            "ILOSTAT rate; every fallback count is recorded."
        ),
        "city_sampling_semantics": (
            "IPF calibrates current country by WUP city population-size class. "
            "The concrete city is sampled within the selected class in "
            "proportion to its WUP 2025 population."
        ),
        "physical_condition_semantics": physical_audit["single_label_semantics"],
        "mental_condition_semantics": mental_audit["single_label_semantics"],
    }

    args.target.write_text(
        json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    args.city_catalog.parent.mkdir(parents=True, exist_ok=True)
    args.city_catalog.write_text(
        json.dumps(city_catalog, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    render_values_doc(data, args.values_doc)
    audit = {
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "ihme_gbd_input": {
            "path": str(args.ihme_gbd.resolve()),
            "size_bytes": args.ihme_gbd.stat().st_size,
            "sha256": hashlib.sha256(args.ihme_gbd.read_bytes()).hexdigest(),
        },
        "source_urls": {
            "wpp_2024_age_sex": WPP_URL,
            "wup_2025_city_population": WUP_CITY_URL,
            "ims_2020_destination_origin": IMS_URL,
            "undata_marital_status_table_23": "https://data.un.org/Data.aspx?d=POP&f=tableCode%3A23",
            "undata_ethnicity_table_26": "https://data.un.org/Data.aspx?d=POP&f=tableCode%3A26",
            "world_bank_pip_api": "https://api.worldbank.org/pip/v1",
            "oecd_relative_income_classification": OECD_MIDDLE_INCOME_URL,
            "ilostat_employment_by_occupation_2025": ILO_OCCUPATION_URL,
            "ilostat_employment_to_population_2025": ILO_EMPLOYMENT_RATE_URL,
            "ilostat_employment_rate_by_age_observed": (
                ILO_EMPLOYMENT_RATE_BY_AGE_URL
            ),
            "ilostat_working_age_population_by_age_education": (
                ILO_WORKING_AGE_AGE_EDUCATION_URL
            ),
            "ilostat_employment_by_age_education": (
                ILO_EMPLOYMENT_AGE_EDUCATION_URL
            ),
            "ilostat_employment_rate_by_education": (
                ILO_EMPLOYMENT_RATE_BY_EDUCATION_URL
            ),
            "ilostat_occupation_by_education": (
                ILO_OCCUPATION_EDUCATION_URL
            ),
            "ilostat_occupation_by_age_stage": (
                ILO_AGE_OCCUPATION_URL
            ),
            "ilostat_employment_rate_by_age_marital_status": (
                ILO_AGE_MARITAL_EMPLOYMENT_RATE_URL
            ),
            "world_bank_global_welfare_profile_2022": (
                WORLD_BANK_WELFARE_PROFILE_URL
            ),
            "pew_religious_composition_2020": PEW_RELIGION_URL,
            "ihme_gbd_2023_results_tool": (
                "https://vizhub.healthdata.org/gbd-results/"
            ),
            "statistics_canada_population_group_2021": (
                "https://www12.statcan.gc.ca/census-recensement/2021/dp-pd/"
                "prof/details/page.cfm?DGUIDlist=2021A000011124"
            ),
            "inegi_census_2020_indigenous": (
                "https://www.inegi.org.mx/contenidos/productos/prod_serv/"
                "contenidos/espanol/bvinegi/productos/nueva_estruc/702825198060.pdf"
            ),
        },
        "migration": migration_audit,
        "migration_status_by_sex": migration_status_audit,
        "education": education_audit,
        "ilo_joint_constraints": ilo_joint_audit,
        "ilo_age_stage_occupation": age_stage_occupation_audit,
        "ilo_relationship_employment": relationship_employment_audit,
        "education_income_joint": education_income_audit,
        "age_income_joint": age_income_audit,
        "religion": religion_audit,
        "income": income_audit,
        "relationship_source_years": marital_years,
        "ethnicity_conditional_distributions": ethnicity_distributions,
        "occupation": occupation_audit,
        "physical_condition": physical_audit,
        "mental_condition": mental_audit,
    }
    args.audit.parent.mkdir(parents=True, exist_ok=True)
    args.audit.write_text(
        json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
